#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""根据演示截图生成 Word 版说明文档。"""
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(ROOT, 'output', 'playwright', 'screenshots')
DOCS = os.path.join(ROOT, 'docs')


def add_image(doc, filename, caption):
    path = os.path.join(SHOTS, filename)
    if not os.path.exists(path):
        return
    doc.add_picture(path, width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_manual():
    doc = Document()
    doc.styles['Normal'].font.name = '微软雅黑'
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.add_heading('错题本知识溯源整理助手', 0)
    doc.add_heading('软件使用说明', 1)
    doc.add_paragraph('版本：v3.2')
    doc.add_paragraph('更新日期：2026-09-24')
    doc.add_paragraph('演示账号：demo    密码：demo123')

    doc.add_heading('1. 登录系统', 1)
    doc.add_paragraph('Windows 推荐双击“启动错题本.bat”，脚本会自动定位 Python、启动服务并打开浏览器。也可以运行 python app.py 后访问 http://localhost:5000。首次进入需要登录，账号和密码只能使用英文字母和数字。')
    add_image(doc, '00_login.png', '图 1 登录界面')

    doc.add_heading('2. 首页概览', 1)
    doc.add_paragraph('首页显示总错题数、未掌握数量、薄弱知识点 TOP3 和最近错题。左侧为垂直导航栏，包含全部功能入口。')
    add_image(doc, '01_dashboard.png', '图 2 首页概览')

    doc.add_heading('3. 错题录入', 1)
    doc.add_paragraph('支持手动录入、拍照 OCR、图片选择和 PDF/Word 导入。题干支持 LaTeX 公式实时渲染，保存前会二次确认。')
    add_image(doc, '02_input.png', '图 3 错题录入')

    doc.add_heading('4. 错题管理', 1)
    doc.add_paragraph('错题库支持按学科、知识点、状态和关键词筛选，支持多选后批量关联知识点、修改状态和删除。错题列表采用分页加载，题目较多时不会一次加载全部数据。')
    add_image(doc, '03_library.png', '图 4 错题库与批量操作')
    doc.add_paragraph('点击题目标题查看详情，可以进行编辑、知识溯源和查看知识导图。')
    add_image(doc, '04_question_detail.png', '图 5 错题详情')

    doc.add_heading('5. 知识溯源思维导图', 1)
    doc.add_paragraph('AI 分析错题后给出错误表象、直接知识点、前置知识点和根本原因。用户勾选后生成独立思维导图，不会写入知识树。')
    add_image(doc, '05_mindmap.png', '图 6 知识溯源思维导图')

    doc.add_heading('6. 复习模式', 1)
    doc.add_paragraph('支持今日推荐、自动生成今日复习计划、自定义日期计划、完成/跳过计划和逾期提醒。可以开启浏览器提醒、设置每日提醒时间，也可以导出 .ics 文件加入手机日历。')
    add_image(doc, '06_review.png', '图 7 复习模式')

    doc.add_heading('7. 薄弱分析', 1)
    doc.add_paragraph('展示薄弱知识点 TOP-N、累计错误、近 7 天错误、30 天趋势和复习优先级。薄弱指数先按当前候选范围内的最高错题数、最高累计错误次数、最高时间衰减值和最高近 7 天错误次数归一化，再按 0.4、0.3、0.2、0.1 加权并乘以 10，得到 0 到 10 的相对分数。没有复习记录时，系统使用错题创建时间估算时间衰减。')
    add_image(doc, '07_weak.png', '图 8 薄弱分析与趋势图')

    doc.add_heading('8. 知识图谱', 1)
    doc.add_paragraph('知识图谱以学科为中心向外发散，支持分层导图、树图和旭日图。点击节点可以查看和编辑知识点说明、知识库内容和学习资源。')
    add_image(doc, '08_graph_layered.png', '图 9 知识图谱分层图')
    add_image(doc, '09_graph_tree.png', '图 10 知识图谱树图')
    add_image(doc, '10_graph_sunburst.png', '图 11 知识图谱旭日图')

    doc.add_heading('9. 统计报表', 1)
    doc.add_paragraph('统计报表展示错题总数、知识点数、近 30 天新增、近 30 天仍错、复习正确率、待复习计划，以及状态分布图和薄弱排行图。支持一键导出 Word 和 PDF。')
    add_image(doc, '11_report.png', '图 12 统计报表')

    doc.add_heading('10. 回收站', 1)
    doc.add_paragraph('删除错题后进入回收站，可以恢复或彻底删除。彻底删除后无法恢复。')
    add_image(doc, '12_trash.png', '图 13 回收站')

    doc.add_heading('11. 学科设置', 1)
    doc.add_paragraph('支持学科管理和知识点树搭建。知识点树逐层展开，支持搜索定位、添加子级和删除子树。')
    add_image(doc, '13_settings.png', '图 14 学科设置')

    doc.add_heading('12. 账号设置与数据备份', 1)
    doc.add_paragraph('账号设置支持改名、修改密码、重置密码、删除账号，以及导出/导入单一数据包。')
    add_image(doc, '14_account.png', '图 15 账号设置')

    doc.add_heading('13. 模型与 OCR 配置', 1)
    doc.add_paragraph('配置大模型 API Key、模型名称、知识点推荐阈值，以及 OCR 语言、方向识别和置信度阈值。')
    add_image(doc, '15_model_ocr.png', '图 16 模型与 OCR 配置')

    doc.add_heading('14. PWA 与离线查看', 1)
    doc.add_paragraph('在 localhost 或 HTTPS 环境可将系统添加到手机主屏幕。断网时可以查看已缓存的错题和首页数据，退出账号时会清理当前用户的离线缓存。局域网普通 HTTP 环境可能无法使用完整安装和系统通知，此时可使用 .ics 日历导出。')

    doc.add_heading('15. 性能与流式输出', 1)
    doc.add_paragraph('错题库采用后端分页，首页只加载最近错题和汇总数量，数据库查询增加索引。概念回顾、知识点精要和单题 AI 解答使用流式输出，内容逐步显示。')

    doc.add_heading('16. 使用限制', 1)
    doc.add_paragraph('拍照识别需要本机安装 PaddleOCR 和 PaddlePaddle；AI 推荐、知识溯源和 AI 学习内容需要配置大模型 API Key。PWA 安装与系统通知需要 localhost 或 HTTPS，浏览器通知只在页面打开时检查。回收站中的错题不会自动过期删除，需要手动彻底删除。')

    out = os.path.join(DOCS, '软件使用说明.docx')
    doc.save(out)
    return out


def build_modules():
    doc = Document()
    doc.styles['Normal'].font.name = '微软雅黑'
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.add_heading('错题本知识溯源整理助手', 0)
    doc.add_heading('功能模块说明', 1)
    doc.add_paragraph('版本：v3.2')
    doc.add_paragraph('更新日期：2026-09-24')
    doc.add_paragraph('实现基线：以当前代码为准，共 13 张数据表、66 个路由路径（78 个 URL/方法绑定）。')

    sections = [
        ('1. 系统概述', '系统面向个人学习者，围绕“错题录入—知识点识别—知识溯源—薄弱分析—复习计划—掌握反馈”构建完整闭环。'),
        ('2. 技术架构', '前端使用 Vue 3、Axios、ECharts、KaTeX；后端使用 Flask、SQLAlchemy；数据库为 SQLite；OCR 使用 PaddleOCR；文档解析使用 python-docx 和 PyMuPDF；大模型支持 DeepSeek 和通义千问。'),
        ('3. 用户账号模块', '提供注册、登录、退出、改名、修改密码、重置密码、删除账号。用户名和密码仅允许英文字母与数字，不同用户数据完全隔离。'),
        ('4. 错题录入模块', '支持手动录入、拍照 OCR、图片上传以及 PDF/Word 试卷导入。导入时自动分题并预览，用户确认后批量入库。'),
        ('5. 错题管理模块', '提供错题增删改查、按学科/知识点/状态筛选、批量关联知识点、批量修改状态和批量删除。删除进入回收站，可恢复或彻底删除。'),
        ('6. 知识树模块', '以学科为根节点，支持多级知识点、逐层展开、搜索定位、添加子级和删除子树。'),
        ('7. 知识溯源模块', '调用大模型分析错误表象、直接知识点、前置知识点和根本原因。用户确认后生成独立思维导图，不污染知识树。用户反馈会影响后续推荐权重。'),
        ('8. 薄弱分析模块', '薄弱指数先按当前候选范围内的最高错题数、最高累计错误次数、最高时间衰减值和最高近 7 天错误次数归一化，再按 0.4、0.3、0.2、0.1 加权并乘以 10，得到 0 到 10 的相对分数。没有复习记录时，使用错题创建时间估算时间衰减。该模块同时提供 30 天趋势和复习优先级。'),
        ('9. 复习模式模块', '支持今日推荐、自动生成今日计划、自定义计划、完成/跳过计划、待复习角标和逾期提醒。'),
        ('10. 知识库模块', '知识点详情中支持知识点说明、知识点精要、概念回顾、典型例题、变式练习等内容的增删改查。'),
        ('11. 学习资源模块', '支持教材页码、网课链接、笔记内容和附件上传。AI 推荐提供 B站搜索链接和复习笔记，不虚构教材页码和 BV 号。'),
        ('12. 知识图谱模块', '提供分层导图、树图和旭日图三种视图，点击节点查看详情，并可调用 AI 精简知识点。'),
        ('13. 统计报表模块', '展示错题总数、知识点数、近 30 天新增、近 30 天仍错、复习正确率、待复习计划，以及状态分布和薄弱排行图表。'),
        ('14. 数据备份模块', '导出单一 ZIP 数据包，内部包含 JSON 和 SQLite 备份；导入时只需选择该数据包，系统自动恢复全部用户数据。'),
        ('15. 系统配置模块', '支持大模型 API Key、模型名称、知识点推荐阈值，以及 OCR 语言、方向识别和置信度阈值配置。'),
        ('16. PWA 移动端模块', '支持添加到主屏幕、离线查看缓存数据和用户级缓存隔离；在 localhost 或 HTTPS 环境下体验最佳。'),
        ('17. 报表导出模块', '统计报表、薄弱分析和复习计划支持导出 Word 与 PDF。'),
        ('18. 性能优化模块', '错题库分页、首页轻量查询、SQLite 索引和 AI 流式输出。'),
    ]
    for title, body in sections:
        doc.add_heading(title, 2)
        doc.add_paragraph(body)

    doc.add_heading('19. 数据库主要数据表', 2)
    doc.add_paragraph('当前代码共创建 13 张数据表：')
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = '表名'
    table.rows[0].cells[1].text = '用途'
    rows = [
        ('users', '用户账号'), ('subjects', '学科'), ('knowledge_points', '知识点树'),
        ('questions', '错题'), ('question_kp', '错题-知识点关联'), ('review_logs', '复习记录'),
        ('review_plans', '复习计划'), ('ai_materials', '知识库内容'),
        ('learning_resources', '学习资源'), ('knowledge_traces', '知识溯源记录'),
        ('mind_maps', '知识溯源思维导图'), ('kp_feedback', '推荐反馈'), ('settings', '用户配置')
    ]
    for name, use in rows:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = use

    doc.add_paragraph('questions.deleted_at 是软删除标记。字段为空表示正常错题，非空表示已进入回收站；当前版本不自动按时间清理回收站。')

    doc.add_heading('20. 演示数据', 2)
    doc.add_paragraph('运行 python tools/seed_demo.py 可生成演示账号 demo / demo123，以及 3 个学科、17 个知识点、12 道错题、30 天复习记录、复习计划、知识库内容、学习资源和知识溯源导图。')

    doc.add_heading('21. 实现边界与已知限制', 2)
    doc.add_paragraph('OCR 依赖本机 PaddleOCR 和 PaddlePaddle；大模型功能依赖用户配置的 API Key。PWA 完整安装和系统通知需要 localhost 或 HTTPS，浏览器通知只在页面打开时检查。SQLite 适合单机个人使用，不适合多机或高并发部署。上传接口尚未统一增加扩展名白名单、文件内容校验和请求体大小限制，部署到公网前应补充这些保护。')

    out = os.path.join(DOCS, '功能模块说明.docx')
    doc.save(out)
    return out


if __name__ == '__main__':
    print(build_manual())
    print(build_modules())
