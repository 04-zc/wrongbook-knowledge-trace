#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""test1 账号全功能自动测试。"""
import io
import json
import os
import sys
import time
from datetime import date, datetime, timedelta

import docx
import fitz
import requests
from PIL import Image, ImageDraw, ImageFont

BASE = 'http://localhost:5000'
RESULTS = []


def check(name, condition, detail=''):
    RESULTS.append((name, bool(condition), detail))
    print(('PASS' if condition else 'FAIL'), name, detail[:120])


def login(username, password):
    s = requests.Session()
    r = s.post(BASE + '/api/auth/login', json={'username': username, 'password': password})
    return s, r


def main():
    s, r = login('test1', '123456')
    check('test1 登录', r.status_code == 200, r.text[:100])
    if r.status_code != 200:
        return

    for old in s.get(BASE + '/api/subjects').json():
        if old['name'].startswith('自动化测试-'):
            s.delete(BASE + '/api/subjects/' + str(old['id']))

    suffix = str(int(time.time()))
    temp_subject_name = '自动化测试-' + suffix
    sub = s.post(BASE + '/api/subjects', json={'name': temp_subject_name}).json()
    check('创建测试学科', bool(sub.get('id')))

    kp1 = s.post(BASE + '/api/knowledge_points', json={
        'subject_id': sub['id'], 'name': '自动化函数', 'parent_id': None
    }).json()
    kp2 = s.post(BASE + '/api/knowledge_points', json={
        'subject_id': sub['id'], 'name': '自动化导数', 'parent_id': kp1['id']
    }).json()
    tree = s.get(BASE + '/api/kp_tree', params={'subject_id': sub['id']}).json()
    check('知识点树父子层级', len(tree) == 2 and tree[1]['parent_id'] == kp1['id'])

    q1 = s.post(BASE + '/api/questions', json={
        'subject_id': sub['id'], 'title_text': '自动化测试：求 x² 的导数。',
        'correct_answer': '2x', 'user_answer': 'x',
        'analysis': '使用幂函数求导公式。',
        'knowledge_points': [{'id': kp2['id'], 'name': kp2['name']}]
    }).json()
    q2 = s.post(BASE + '/api/questions', json={
        'subject_id': sub['id'], 'title_text': '自动化测试：求函数定义域。'
    }).json()
    check('错题录入', bool(q1.get('id') and q2.get('id')))
    check('按知识点筛选错题',
          len(s.get(BASE + '/api/questions', params={'kp_id': kp2['id']}).json()) == 1)
    check('编辑错题',
          s.put(BASE + '/api/questions/' + str(q2['id']),
                json={'title_text': '自动化测试：求函数定义域（已修改）'}).status_code == 200)

    check('批量修改状态',
          s.post(BASE + '/api/questions/batch', json={
              'action': 'status', 'ids': [q1['id'], q2['id']], 'status': 2
          }).status_code == 200)
    check('批量关联知识点',
          s.post(BASE + '/api/questions/batch', json={
              'action': 'bind_kp', 'ids': [q2['id']], 'kp_ids': [kp1['id']]
          }).status_code == 200)

    s.post(BASE + '/api/questions/' + str(q1['id']) + '/review', json={'result': 0})
    s.post(BASE + '/api/questions/' + str(q1['id']) + '/review', json={'result': 0})
    weak = s.get(BASE + '/api/weak_points', params={'subject_id': sub['id']}).json()
    weak_item = [x for x in weak if x['id'] == kp2['id']][0]
    check('薄弱指数完整计算',
          weak_item['error_count'] >= 2 and weak_item['recent_errors'] >= 2 and weak_item['weak_index'] > 0,
          str({k: weak_item[k] for k in ['error_count', 'recent_errors', 'decay_score', 'weak_index']}))
    check('30 天趋势接口',
          len(s.get(BASE + '/api/analytics/trend', params={'subject_id': sub['id']}).json()) == 30)
    check('复习优先级接口',
          isinstance(s.get(BASE + '/api/analytics/priority', params={'subject_id': sub['id']}).json(), list))

    today = date.today().isoformat()
    plan = s.post(BASE + '/api/review/plans', json={
        'question_id': q1['id'], 'plan_date': today, 'note': '自动化测试'
    }).json()
    plan_data = s.get(BASE + '/api/review/plans').json()
    check('创建复习计划', plan_data['counts']['today'] >= 1, str(plan_data['counts']))
    check('完成复习计划',
          s.post(BASE + '/api/review/plans/' + str(plan['id']) + '/complete',
                 json={'result': 1}).status_code == 200)

    material_res = s.post(BASE + '/api/ai/materials', json={
        'kp_id': kp2['id'], 'mode': 'custom', 'content': '自动化知识内容'
    }).json()
    material = material_res.get('material') or material_res
    check('知识库内容新增',
          len(s.get(BASE + '/api/ai/materials', params={'kp_id': kp2['id']}).json()) >= 1)
    check('知识库内容修改',
          s.post(BASE + '/api/ai/materials/' + str(material['id']),
                 json={'content': '自动化知识内容（修改）'}).status_code == 200)

    resource = s.post(BASE + '/api/knowledge_points/' + str(kp2['id']) + '/resources', json={
        'type': 'textbook', 'title': '自动化教材', 'page': '10', 'content': '自动化笔记'
    }).json()
    check('学习资源新增',
          len(s.get(BASE + '/api/knowledge_points/' + str(kp2['id']) + '/resources').json()) >= 1)
    check('学习资源修改',
          s.post(BASE + '/api/resources/' + str(resource['id']),
                 json={'page': '11'}).status_code == 200)

    trace = s.post(BASE + '/api/questions/' + str(q1['id']) + '/trace/confirm', json={
        'symptom': '求导错误', 'root_cause': '公式不熟',
        'points': [{'name': '自动化溯源点', 'reason': '直接考查', 'group': 'direct_points'}]
    })
    check('知识溯源导图生成',
          trace.status_code == 200 and len(trace.json()['mindmap']['nodes']) >= 3)
    check('导图保存',
          s.get(BASE + '/api/questions/' + str(q1['id']) + '/mindmap').status_code == 200)
    check('溯源未写入知识树',
          all(k['name'] != '自动化溯源点' for k in s.get(
              BASE + '/api/knowledge_points', params={'subject_id': sub['id']}).json()))

    rec = s.post(BASE + '/api/questions/' + str(q1['id']) + '/recommend_kp')
    check('AI 知识点推荐接口', rec.status_code == 200, rec.text[:120])

    ai_checks = [
        ('AI 概念回顾', {'kp_id': kp2['id'], 'mode': 'concept'}, 'content'),
        ('AI 典型例题', {'kp_id': kp2['id'], 'mode': 'typical'}, 'questions'),
        ('AI 变式练习', {'kp_id': kp2['id'], 'mode': 'variant'}, 'questions'),
        ('AI 知识点精要', {'kp_id': kp2['id'], 'mode': 'summary'}, 'content'),
        ('AI 学习资源推荐', {'kp_id': kp2['id'], 'mode': 'resources'}, 'resources'),
        ('AI 单题解答', {'question_id': q1['id'], 'mode': 'explain'}, 'content'),
    ]
    for name, payload, key in ai_checks:
        res = s.post(BASE + '/api/ai/generate', json=payload, timeout=120)
        ok = res.status_code == 200 and bool(res.json().get(key))
        check(name, ok, res.text[:150])

    # OCR
    img = Image.new('RGB', (480, 110), 'white')
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype('C:/Windows/Fonts/simhei.ttf', 26)
    draw.text((12, 34), '已知函数f(x)=2x+1，求导', fill='black', font=font)
    img_bytes = io.BytesIO()
    img.save(img_bytes, format='PNG')
    ocr = s.post(BASE + '/api/ocr',
                 files={'file': ('ocr.png', img_bytes.getvalue(), 'image/png')},
                 timeout=120)
    check('OCR 图片识别', ocr.status_code == 200 and bool(ocr.json().get('text')), ocr.text[:120])

    # Word / PDF
    d = docx.Document()
    d.add_paragraph('1. Word 自动测试第一题')
    d.add_paragraph('2. Word 自动测试第二题')
    wb = io.BytesIO()
    d.save(wb)
    word = s.post(BASE + '/api/import/document',
                  files={'file': ('auto.docx', wb.getvalue(),
                                  'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    check('Word 导入解析', word.status_code == 200 and word.json().get('count', 0) >= 2,
          word.text[:120])
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((50, 70), '1. PDF auto test question one', fontsize=12)
    page.insert_text((50, 100), '2. PDF auto test question two', fontsize=12)
    pdf_res = s.post(BASE + '/api/import/document',
                     files={'file': ('auto.pdf', pdf.tobytes(), 'application/pdf')})
    check('PDF 导入解析', pdf_res.status_code == 200, pdf_res.text[:120])

    check('知识图谱接口',
          s.get(BASE + '/api/knowledge_graph', params={'subject_id': sub['id']}).status_code == 200)
    check('统计报表接口',
          s.get(BASE + '/api/reports/summary', params={'subject_id': sub['id']}).status_code == 200)

    # 回收站
    trash_q = s.post(BASE + '/api/questions', json={
        'subject_id': sub['id'], 'title_text': '回收站自动测试'
    }).json()
    s.delete(BASE + '/api/questions/' + str(trash_q['id']))
    check('移入回收站',
          any(x['id'] == trash_q['id'] for x in s.get(BASE + '/api/questions/trash').json()))
    s.post(BASE + '/api/questions/' + str(trash_q['id']) + '/restore')
    check('回收站恢复',
          any(x['id'] == trash_q['id'] for x in s.get(BASE + '/api/questions').json()))
    s.delete(BASE + '/api/questions/' + str(trash_q['id']))
    s.delete(BASE + '/api/questions/' + str(trash_q['id']) + '/purge')
    check('彻底删除',
          not any(x['id'] == trash_q['id'] for x in s.get(BASE + '/api/questions/trash').json()))

    # 数据包导出与恢复
    package = s.get(BASE + '/api/backup/export')
    check('导出单一数据包',
          package.status_code == 200 and package.content[:2] == b'PK')
    temp_user = 'e2euser' + suffix[-5:]
    ts = requests.Session()
    ts.post(BASE + '/api/auth/register', json={'username': temp_user, 'password': 'abc123'})
    imported = ts.post(BASE + '/api/backup/import',
                       files={'file': ('backup.zip', package.content, 'application/zip')})
    check('导入数据包', imported.status_code == 200 and imported.json().get('ok'), imported.text[:120])
    restored_weak = ts.get(BASE + '/api/weak_points').json()
    restored_errors = max([x.get('error_count', 0) for x in restored_weak] or [0])
    check('恢复复习记录和薄弱指数', restored_errors >= 2, 'error_count=' + str(restored_errors))
    ts.delete(BASE + '/api/auth/account', json={'password': 'abc123'})

    # 账号功能使用独立测试账号，避免修改 test1
    account_user = 'e2eacc' + suffix[-5:]
    ac = requests.Session()
    ac.post(BASE + '/api/auth/register', json={'username': account_user, 'password': 'abc123'})
    check('账号修改密码',
          ac.post(BASE + '/api/auth/change_password',
                  json={'old_password': 'abc123', 'new_password': 'abc456'}).status_code == 200)
    renamed = account_user + 'x'
    check('账号改名',
          ac.post(BASE + '/api/auth/rename',
                  json={'new_username': renamed, 'password': 'abc456'}).status_code == 200)
    check('账号重置密码',
          requests.post(BASE + '/api/auth/reset_password',
                        json={'username': renamed, 'new_password': 'abc789'}).status_code == 200)
    ac2 = requests.Session()
    ac2.post(BASE + '/api/auth/login', json={'username': renamed, 'password': 'abc789'})
    check('账号删除',
          ac2.delete(BASE + '/api/auth/account', json={'password': 'abc789'}).status_code == 200)

    # 清理测试学科
    s.delete(BASE + '/api/subjects/' + str(sub['id']))

    passed = sum(1 for _, ok, _ in RESULTS if ok)
    report = [
        '# test1 全功能测试报告',
        '',
        f'- 测试时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}',
        f'- 结果：{passed}/{len(RESULTS)} 项通过',
        '',
        '| 功能 | 结果 | 说明 |',
        '|---|---|---|',
    ]
    for name, ok, detail in RESULTS:
        report.append(f'| {name} | {"通过" if ok else "失败"} | {detail.replace("|", " ")[:120]} |')
    os.makedirs('output', exist_ok=True)
    with open('output/test_report_test1.md', 'w', encoding='utf-8') as f:
        f.write('\n'.join(report))
    print('PASS', passed, '/', len(RESULTS))
    print('report: output/test_report_test1.md')


if __name__ == '__main__':
    main()
