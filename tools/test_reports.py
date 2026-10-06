# -*- coding: utf-8 -*-
"""统计报表逐页核对。

生成 Word 和 PDF 报表后逐页读取文本，检查标题、统计值和表格内容
是否真的写进文件，而不是只检查接口状态码。

运行：python tools/test_reports.py
"""
import io
import os
import shutil
import sys
import tempfile
from datetime import datetime, timedelta

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import db

TMP = tempfile.mkdtemp()
db.DB_PATH = os.path.join(TMP, 'reports.db')
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
db.engine = create_engine('sqlite:///' + db.DB_PATH)
db.SessionLocal = sessionmaker(bind=db.engine)
db.Base.metadata.create_all(db.engine)

import docx
import fitz
import report_export
import app as webapp

webapp.app.config['TESTING'] = True
webapp.app.config['CSRF_ENABLED'] = False
client = webapp.app.test_client()

PASS, FAIL = [], []


def chk(name, cond, detail=''):
    (PASS if cond else FAIL).append(
        name + ((' :: ' + str(detail)) if detail != '' and not cond else '')
    )


client.post('/api/auth/register', json={'username': 'reportuser', 'password': 'pass11'})
db.set_current_user(1)
subject = db.add_subject('数学')
kp_function = db.add_kp(subject['id'], '函数')
kp_sequence = db.add_kp(subject['id'], '数列')

questions = []
for i in range(3):
    questions.append(db.add_question({
        'subject_id': subject['id'],
        'title_text': '函数题 %d' % i,
        'status': 0,
        'knowledge_points': [{'id': kp_function['id']}],
    }))
questions.append(db.add_question({
    'subject_id': subject['id'],
    'title_text': '数列题',
    'status': 1,
    'knowledge_points': [{'id': kp_sequence['id']}],
}))
db.add_review(questions[0]['id'], 0, '仍然做错')
db.add_review(questions[1]['id'], 1, '订正完成')
today = datetime.utcnow().date()
db.add_review_plan(questions[0]['id'], today.isoformat(), '复习函数')
db.add_review_plan(questions[3]['id'], (today + timedelta(days=1)).isoformat(), '复习数列')

data = report_export.collect_report()
chk('统计汇总错题数为 4', data['totals']['questions'] == 4, data['totals']['questions'])
chk('统计汇总知识点数为 2', data['totals']['knowledge_points'] == 2,
    data['totals']['knowledge_points'])
chk('复习计划列表为 2 条', len(data['plans']) == 2, data['plans'])
chk('今日到期计划为 1（未来计划不计入待复习）',
    data['totals']['plans']['total'] == 1, data['totals']['plans'])

# ---------- Word：逐段逐表核对 ----------
word_response = client.get('/api/reports/export/word')
chk('Word 导出接口返回 200', word_response.status_code == 200, word_response.status_code)
chk('Word 文件头正确', word_response.data[:2] == b'PK', word_response.data[:8])

word_doc = docx.Document(io.BytesIO(word_response.data))
word_paragraphs = '\n'.join(p.text for p in word_doc.paragraphs)
word_tables = []
for table in word_doc.tables:
    rows = []
    for row in table.rows:
        rows.append([cell.text for cell in row.cells])
    word_tables.append(rows)

for heading in ['一、总体统计', '二、薄弱知识点 TOP-N', '三、复习计划', '四、近 30 天趋势']:
    chk('Word 包含标题：' + heading, heading in word_paragraphs)
for expected in ['错题总数：4', '知识点数量：2', '待复习计划：1', '函数题 0']:
    chk('Word 包含内容：' + expected, expected in word_paragraphs or any(
        expected in cell for table in word_tables for row in table for cell in row
    ))
chk('Word 生成三个数据表格', len(word_tables) == 3, len(word_tables))
chk('Word 薄弱表格含知识点行',
    any(row and row[0] == '函数' for row in (word_tables[0] if word_tables else [])),
    word_tables[0] if word_tables else None)
chk('Word 复习计划表格包含两天计划',
    len(word_tables) >= 2 and len(word_tables[1]) == 3,
    word_tables[1] if len(word_tables) >= 2 else None)

# ---------- PDF：逐页核对，每页都必须有文字 ----------
pdf_response = client.get('/api/reports/export/pdf')
chk('PDF 导出接口返回 200', pdf_response.status_code == 200, pdf_response.status_code)
chk('PDF 文件头正确', pdf_response.data[:5] == b'%PDF-', pdf_response.data[:8])

pdf = fitz.open(stream=pdf_response.data, filetype='pdf')
page_texts = [page.get_text() for page in pdf]
pdf.close()
chk('PDF 至少一页', len(page_texts) >= 1, len(page_texts))
chk('PDF 每页都不是空白', all(text.strip() for text in page_texts),
    [len(text.strip()) for text in page_texts])
pdf_all_text = '\n'.join(page_texts)
for heading in ['一、总体统计', '二、薄弱知识点 TOP-N', '三、待复习计划']:
    chk('PDF 包含标题：' + heading, heading in pdf_all_text)
chk('PDF 包含错题总数', '错题总数' in pdf_all_text and '4' in pdf_all_text)
chk('PDF 包含薄弱知识点', '函数' in pdf_all_text, pdf_all_text[:200])

# ---------- 汇总接口与报表数据一致 ----------
summary = client.get('/api/reports/summary')
summary_data = summary.get_json() if summary.status_code == 200 else {}
chk('汇总接口错题数与报表一致',
    summary.status_code == 200 and summary_data.get('totals', {}).get('questions') == 4,
    summary_data.get('totals'))

print('=== PASS (%d) ===' % len(PASS))
for name in PASS:
    print('  ok  ' + name)
if FAIL:
    print('=== FAIL (%d) ===' % len(FAIL))
    for name in FAIL:
        print('  XX  ' + name)
shutil.rmtree(TMP, ignore_errors=True)
sys.exit(1 if FAIL else 0)
