#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将根目录需求说明书更新到当前代码基线，同时保留原有文档结构。"""

import ast
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRS_PATH = os.path.join(ROOT, '错题本知识溯源整理助手_软件需求规格说明书.docx')


def set_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def find_paragraph(doc, prefix):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise KeyError(f'paragraph not found: {prefix}')


def replace_paragraph(doc, prefix, text):
    paragraph = find_paragraph(doc, prefix)
    set_text(paragraph, text)
    return paragraph


def find_table(doc, header_cell):
    for table in doc.tables:
        if table.rows and table.rows[0].cells and table.rows[0].cells[0].text.strip() == header_cell:
            return table
    raise KeyError(f'table not found: {header_cell}')


def fill_table(table, headers, rows):
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def next_table_after_heading(doc, heading_prefix):
    heading = find_paragraph(doc, heading_prefix)
    sibling = heading._p.getnext()
    while sibling is not None:
        if sibling.tag == qn('w:tbl'):
            return Table(sibling, doc)
        if sibling.tag == qn('w:p'):
            paragraph = Paragraph(sibling, doc)
            if paragraph.style.name.startswith('Heading'):
                break
        sibling = sibling.getnext()
    raise KeyError(f'status table not found after: {heading_prefix}')


def set_status(doc, heading_prefix, priority, status):
    table = next_table_after_heading(doc, heading_prefix)
    table.rows[0].cells[1].text = priority
    table.rows[0].cells[3].text = status


def add_before(doc, anchor, text, style=None):
    paragraph = doc.add_paragraph(text, style=style)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def add_table_before(doc, anchor, headers, rows, repeat_header=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    fill_table(table, headers, rows)
    if repeat_header:
        set_repeat_table_header(table.rows[0])
    anchor._p.addprevious(table._tbl)
    return table


def route_bindings():
    source = open(os.path.join(ROOT, 'app.py'), encoding='utf-8').read()
    tree = ast.parse(source)
    paths = {}
    for node in tree.body:
        if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            continue
        for decorator in node.decorator_list:
            if not (
                isinstance(decorator, ast.Call)
                and isinstance(decorator.func, ast.Attribute)
                and isinstance(decorator.func.value, ast.Name)
                and decorator.func.value.id == 'app'
            ):
                continue
            if decorator.func.attr not in ('route', 'get', 'post', 'put', 'delete', 'patch'):
                continue
            path = decorator.args[0].value
            method_values = []
            for keyword in decorator.keywords:
                if keyword.arg == 'methods' and isinstance(keyword.value, ast.List):
                    method_values = [item.value.upper() for item in keyword.value.elts]
            if not method_values:
                method_values = [decorator.func.attr.upper() if decorator.func.attr != 'route' else 'GET']
            paths.setdefault(path, set()).update(method_values)
    return [(path, sorted(methods)) for path, methods in sorted(paths.items())]


def main():
    doc = Document(SRS_PATH)

    cover = doc.tables[0]
    cover.rows[1].cells[1].text = 'v3.2（代码对齐版）'
    cover.rows[2].cells[1].text = '2026 年 9 月 24 日'
    cover.rows[3].cells[1].text = 'Flask + SQLAlchemy + SQLite + Vue3 + Axios + ECharts + KaTeX'

    replace_paragraph(
        doc,
        '本文档是',
        '本文档是"错题本知识溯源整理助手"项目的软件需求规格说明书。系统面向个人学习者，'
        '以错题为数据入口，通过"大模型结构化解析 + 本地规则兜底"推荐知识点，采用半自动确认机制建立'
        '错题与个人知识树的可靠关联，并通过薄弱指数、复习计划、知识溯源导图和统计报表形成完整学习闭环。'
        '本文档以 2026-09-24 当前代码为基线，覆盖 10 个功能模块、44 项需求、13 张数据表、66 个路由路径'
        '（78 个 URL/方法绑定），并逐项标注实现状态，作为设计、开发、测试与验收依据。'
    )

    status_definition = find_table(doc, '标识')
    fill_table(
        status_definition,
        ['标识', '含义'],
        [
            ('已实现', '当前代码已提供对应后端能力；涉及界面或外部依赖时，按需求说明中的条件使用'),
            ('部分实现', '后端接口已实现，但缺少统一配置入口、前端能力或完整流程'),
            ('规划中', '尚未开发，本文档给出可落地的规格供后续迭代'),
        ],
    )

    toc_appendix = find_paragraph(doc, '附录 C')
    appendix_d = doc.add_paragraph('附录 D　完整接口索引', style=toc_appendix.style)
    toc_appendix._p.addnext(appendix_d._p)

    set_status(doc, 'FR-AUTH-07', 'P1', '已实现')
    set_status(doc, 'FR-SUBJ-03', 'P0', '已实现')
    set_status(doc, 'FR-INP-06', 'P1', '已实现')
    set_status(doc, 'FR-QM-05', 'P1', '已实现')
    set_status(doc, 'FR-WEAK-05', 'P1', '已实现')
    set_status(doc, 'FR-REV-01', 'P0', '已实现')
    set_status(doc, 'FR-REV-02', 'P1', '已实现')
    set_status(doc, 'FR-SET-03', 'P1', '已实现')

    replacements = {
        '功能描述：提供本地账号列表管理':
            '功能描述：提供当前账号的改名、修改密码、重置密码和删除账号能力，各项操作均校验当前登录用户。',
        '新增接口 GET /api/auth/accounts':
            '改名：POST /api/auth/rename，入参 new_username、password；新用户名需符合账号格式且不能与其他账号重复。',
        '修改密码：POST /api/auth/change_password':
            '修改密码：POST /api/auth/change_password，入参 old_password、new_password；校验原密码后更新哈希。',
        '删除账号：DELETE /api/auth/accounts/{username}':
            '重置密码：POST /api/auth/reset_password，入参 username、new_password；用于本地账号密码恢复。',
        '前端入口置于登录页"账号管理"链接':
            '删除账号：DELETE /api/auth/account，入参 password；校验密码后级联删除该用户的业务数据并清空会话。前端入口位于左下角"账号设置"。',
        '验收标准：修改密码后旧密码无法登录':
            '验收标准：改名后旧账号无法登录且新账号可登录；修改或重置密码后旧密码失效、新密码可登录；删除账号后该用户数据清除且其他账号不受影响。',
        '后端 DELETE /api/subjects/{id} 已实现':
            '后端 DELETE /api/subjects/{id} 与前端删除入口均已实现，前端提供二次确认提示。',
        '前端需补齐删除按钮与二次确认弹窗':
            '删除时级联清理该学科下全部错题（含复习记录、知识点关联）与全部知识点；删除后刷新学科、知识点、错题、薄弱分析和统计报表。',
        '删除后刷新学科、知识点、错题、薄弱分析四类数据。':
            '当前删除为物理级联删除，执行后不可恢复。',
        '上传接口 POST /api/import_document':
            '上传接口 POST /api/import/document，支持扩展名 .pdf、.docx、.doc；PDF 使用文本提取或 OCR，Word 使用 python-docx 解析。',
        '解析流程：提取全文':
            '解析流程：提取全文，按题号规则分题，返回题干预览、题号与页码等信息；旧版 .doc 文件可能需要先转换为 .docx。',
        '用户在预览界面勾选错题':
            '前端预览页可勾选、修改并指定学科，点击确认后调用 POST /api/import/save 批量入库，source 字段记录为 import。',
        '确认后批量写入错题库':
            '导入前先解析文档并预览；确认保存时提交 subject_id 与 questions 数组，后端逐题校验题干后写入。',
        '分题识别置信度低的段落标记"需人工确认"':
            '分题结果全部进入人工预览，不直接写库；题干为空或用户取消勾选的题目会被跳过。',
        '单文件上限 20MB，超出提示拒绝。':
            '当前版本尚未统一设置上传大小上限，部署到公网前应补充请求体大小限制、扩展名白名单与内容校验。',
        '功能描述：删除错题（物理删除）':
            '功能描述：删除错题时使用软删除，将 questions.deleted_at 设为当前时间；错题进入回收站，关联记录暂时保留。',
        '验收标准：删除后错题从列表消失':
            '验收标准：删除后错题从错题库、薄弱分析、复习推荐和统计报表中消失；恢复后重新出现；只有彻底删除才清理 question_kp、review_logs 等关联记录。',
        '规格要点：错题删除改为软删除':
            '规格要点：删除采用软删除；回收站提供恢复与彻底删除。当前版本不按时间自动清理，只有用户执行彻底删除才物理移除数据。',
        '功能描述：以知识点为统计单元，将其关联错题规模与当前未掌握程度量化为薄弱指数':
            '功能描述：以知识点为统计单元，综合错题数量、累计错误次数、时间衰减和近 7 天重复错误计算相对薄弱指数。',
        '薄弱指数 = 0.4 × Q':
            '薄弱指数 = (Q/Qmax)×0.4 + (E/Emax)×0.3 + (D/Dmax)×0.2 + (R7/R7max)×0.1，最后乘以 10 并保留两位小数。',
        '结果四舍五入保留两位小数。':
            'Q 为未删除错题数；E 为 review_logs 中 result=0 的累计错误次数；D 为错误时间衰减值之和；R7 为最近 7 天 result=0 的次数。各分项分别除以当前候选集合中的最大值。',
        '示例：某知识点关联错题 5 道':
            '当某知识点在四个分项上都达到当前候选范围最大值时，指数为 10.00；没有复习记录时，时间衰减使用错题 created_at 估算。',
        '验收标准：构造已知 Q、U 的知识点':
            '验收标准：构造固定候选集合与分项数据，接口返回的 weak_index 与归一化公式一致（误差不超过 0.01）。',
        '功能描述：按学科聚合，输出薄弱指数降序排列的 TOP-N 清单':
            '功能描述：按学科聚合，输出薄弱指数降序排列的 TOP-N 清单（薄弱分析页默认 N=20，首页展示 TOP3）。表格列包含知识点、学科、错题数、未掌握数、累计错误、近 7 天重复、时间衰减和薄弱指数。',
        '功能描述：薄弱指数按区间着色':
            '功能描述：薄弱指数按当前阈值着色：小于 4 绿色（#10b981），4 至 7 橙色（#f59e0b），不小于 7 红色（#dc2626）。',
        '验收标准：三个区间的指数分别显示对应颜色。':
            '验收标准：三个区间的指数分别显示对应颜色，颜色计算与 heatColor 规则一致。',
        '规格要点：引入 ECharts':
            '实现规格：知识图谱提供分层导图、树图和旭日图；趋势接口提供近 30 天错误变化；优先级接口给出高中低建议。',
        '验收标准：旭日图层级与知识树一致':
            '验收标准：三种图谱与知识树层级一致；趋势数据与错题创建时间、复习记录一致；优先级建议随薄弱指数变化。',
        '待开发（前端）：错题详情 / 错题库提供':
            '前端已在复习模式提供单题复习、显示答案、标记已掌握或仍需加强，并在提交后刷新薄弱分析。',
        '规格要点：当日复习':
            '实现规格：GET /api/review/today 返回今日建议、推荐知识点与知识树；GET /api/review/plans 返回按逾期、今天、未来分组的复习计划。',
        '自定义计划：按薄弱知识点选择复习范围与频次':
            '计划创建：POST /api/review/plans，入参 question_id、plan_date、note；POST /api/review/plans/generate 可自动生成今日计划。',
        '复习交互：逐题展示题干':
            '复习交互：逐题展示题干，可隐藏答案；完成计划时提交 result，系统同步更新错题状态、复习记录和薄弱分析。',
        '计划数据新增 review_plans 表':
            '复习计划使用 review_plans 表（user_id、question_id、plan_date、status、note、created_at），并支持 DELETE /api/review/plans/{id} 与 GET /api/review/plans/export.ics。',
        '验收标准：当日复习列表按薄弱程度排序':
            '验收标准：今日推荐按薄弱程度排序；计划可完成、跳过、删除；逾期数量正确；导出的 .ics 日历可被手机系统识别。',
        '规格要点：':
            '实现规格：',
        '接口规格（新增）：POST /api/knowledge_points/{id}/concept_review':
            '接口规格：POST /api/ai/generate 或 /api/ai/stream，入参 {kp_id, mode:"concept"}；生成后可调用 /api/ai/save 保存为知识库内容。',
        '接口规格（新增）：GET /api/knowledge_points/{id}/typical_questions':
            '接口规格：GET /api/kp/typical?kp_id={id}&limit=5，返回该知识点关联的典型例题。',
        '接口规格（新增）：POST /api/knowledge_points/{id}/variant_questions':
            '接口规格：POST /api/ai/generate，入参 {kp_id, mode:"variant", count}；返回变式题列表，count 在当前实现中截断为 1~5。',
        '数据规格（新增表 kp_resources）':
            '数据规格：使用 learning_resources 表保存学习资源；接口为 GET/POST /api/knowledge_points/{kid}/resources、POST/DELETE /api/resources/{rid}，附件通过 /api/upload_attachment 上传。',
        '已实现接口清单如下：':
            '以下列出核心接口，完整 66 个路由路径及 HTTP 方法见附录 D。',
        '规划中接口（P1）：账号管理':
            '账号管理、文档导入、知识溯源、复习计划、AI 学习资源、学习资源、数据备份等接口均已在当前代码中实现；具体路径以附录 D 为准。',
        '数据库为 SQLite 单文件（data.db），共 7 张核心表。':
            '数据库为 SQLite 单文件（data.db），当前共 13 张表。核心 7 张为 users、subjects、knowledge_points、questions、question_kp、review_logs、settings；扩展 6 张为 review_plans、ai_materials、knowledge_traces、mind_maps、kp_feedback、learning_resources。',
        '以知识点为统计单元，SQL 聚合获得 Q':
            '算法先按知识点聚合 Q（未删除错题数），再逐知识点读取 review_logs；E 统计 result=0 的累计错误，D 累加 exp(-距今天数/30)，R7 统计最近 7 天 result=0 的次数。没有复习记录时用错题创建时间估算 D。',
        'SELECT kp.id, kp.name,':
            '各分项分别除以当前候选集合中的 Qmax、Emax、Dmax、R7max，按 0.4、0.3、0.2、0.1 加权后乘以 10；分母缺失或为 0 时按 1 处理。',
        '代入公式 weak_index = round(qc×0.4':
            '算法不使用机器学习，权重固定、结果可解释；删除错题不参与统计，恢复后重新参与。',
        '薄弱指数：构造 Q=5、U=3 等样例':
            '薄弱指数：构造包含多个知识点的候选集合，验证分项归一化、加权、四舍五入和无复习记录回退。',
    }

    for prefix, text in replacements.items():
        paragraph = next((p for p in doc.paragraphs if p.text.strip().startswith(prefix)), None)
        if paragraph is not None:
            set_text(paragraph, text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if '管理本地账号（规划中）' in cell.text:
                    set_text(cell.paragraphs[0], '备份导出数据、管理本地账号与回收站')
            if row.cells and row.cells[0].text.strip() == '账号与导入':
                row.cells[2].text = '已完成账号管理、PDF/Word 导入、学科删除和复习前端入口'
                row.cells[3].text = 'FR-AUTH-*、FR-INP-06、FR-SUBJ-03、FR-REV-01/02'

    weak_table = None
    for table in doc.tables:
        if table.rows and len(table.columns) == 3 and table.rows[0].cells[0].text.strip() == '因子':
            weak_table = table
            break
    if weak_table is None:
        raise KeyError('weak factor table not found')
    fill_table(
        weak_table,
        ['因子', '含义', '计算说明'],
        [
            ('Q', '知识点关联的未删除错题数', '历史错题规模'),
            ('E', 'review_logs 中 result=0 的累计次数', '累计错误次数'),
            ('D', 'Σ e^(-距离今天的天数/30)', '时间衰减；无复习记录时用 created_at 回退'),
            ('R7', '最近 7 天内 result=0 的次数', '近期重复错误'),
            ('归一化', '各分项 / 候选集合中的最大值', '最大值缺失或为 0 时按 1 处理'),
        ],
    )

    questions_table = None
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == '字段' and any(
            row.cells[0].text.strip() == 'questions' for row in table.rows
        ):
            questions_table = table
            break
    if questions_table is None:
        for table in doc.tables:
            if table.rows and table.rows[0].cells[0].text.strip() == '字段' and len(table.rows) >= 14:
                texts = [row.cells[0].text.strip() for row in table.rows]
                if 'source' in texts and 'updated_at' in texts:
                    questions_table = table
                    break
    if questions_table is not None and not any(row.cells[0].text.strip() == 'deleted_at' for row in questions_table.rows):
        cells = questions_table.add_row().cells
        values = ('deleted_at', 'DATETIME', '可空，索引', '软删除时间；为空表示正常，非空表示在回收站')
        for cell, value in zip(cells, values):
            cell.text = value

    api_rows = [
        ('API-01', 'GET /api/auth/me', '查询当前登录用户', '—', '{user:{id,username}}', '401 未登录'),
        ('API-02', 'POST /api/auth/register', '注册账号', 'username, password', '{user:{id,username}}', '400 格式非法 / 已存在'),
        ('API-03', 'POST /api/auth/login', '登录', 'username, password', '{user:{id,username}}', '400 账号或密码错误'),
        ('API-04', 'POST /api/auth/logout', '退出登录', '—', '{ok:true}', '—'),
        ('API-05', 'POST /api/auth/rename', '修改账号名', 'new_username, password', '{ok:true,username}', '400 密码或账号错误'),
        ('API-06', 'POST /api/auth/change_password', '修改密码', 'old_password, new_password', '{ok:true}', '400 原密码错误'),
        ('API-07', 'POST /api/auth/reset_password', '重置密码', 'username, new_password', '{ok:true}', '404 账号不存在'),
        ('API-08', 'DELETE /api/auth/account', '删除当前账号', 'password', '{ok:true}', '400 密码错误'),
        ('API-09', 'GET /api/subjects', '学科列表', '—', '学科数组', '401'),
        ('API-10', 'POST /api/subjects', '添加学科', 'name, icon?', '学科对象', '唯一约束冲突'),
        ('API-11', 'DELETE /api/subjects/{sid}', '删除学科及下级数据', '—', '{ok:true}', '401'),
        ('API-12', 'GET /api/knowledge_points', '知识点列表', 'subject_id?', '知识点数组', '401'),
        ('API-13', 'POST /api/knowledge_points', '添加知识点', 'subject_id, name, parent_id?', '知识点对象', '400 父节点 / 学科无效'),
        ('API-14', 'DELETE /api/knowledge_points/{kid}', '删除知识点子树', '—', '{ok:true}', '401'),
        ('API-15', 'GET /api/questions/page', '错题分页列表', 'page, page_size, filters', '分页对象', '401'),
        ('API-16', 'GET /api/questions/{qid}', '错题详情', '—', '错题对象', '404'),
        ('API-17', 'POST /api/questions', '新增错题', '完整表单字段', '错题对象', '400 / 401'),
        ('API-18', 'PUT /api/questions/{qid}', '更新错题', '待修改字段', '错题对象', '404'),
        ('API-19', 'DELETE /api/questions/{qid}', '软删除错题', '—', '{ok:true}', '404'),
        ('API-20', 'POST /api/questions/{qid}/restore', '恢复错题', '—', '{ok:true}', '404'),
        ('API-21', 'DELETE /api/questions/{qid}/purge', '彻底删除错题', '—', '{ok:true}', '404'),
        ('API-22', 'POST /api/questions/{qid}/recommend_kp', 'AI 知识点推荐', '—', '候选知识点数组', '404'),
        ('API-23', 'POST /api/questions/{qid}/confirm_kp', '确认知识点绑定', 'knowledge_points', '{ok:true}', '400 / 404'),
        ('API-24', 'POST /api/questions/{qid}/trace', '生成知识溯源', '—', '溯源结果', '404 / 400'),
        ('API-25', 'POST /api/questions/{qid}/review', '记录复习结果', 'result, note', '{ok:true}', '404'),
        ('API-26', 'GET /api/weak_points', '薄弱知识点 TOP-N', 'subject_id?, top_n', '统计数组', '401'),
        ('API-27', 'GET /api/review/today', '今日复习推荐', 'subject_id?', '建议与推荐知识点', '401'),
        ('API-28', 'GET/POST /api/review/plans', '查询或新增复习计划', 'question_id, plan_date, note', '计划对象 / 列表', '400 / 404'),
        ('API-29', 'POST /api/review/plans/generate', '生成今日计划', 'subject_id?, limit', '{created,plans}', '400'),
        ('API-30', 'GET /api/knowledge_graph', '知识图谱数据', 'subject_id', '{subject,nodes}', '400 / 404'),
        ('API-31', 'GET /api/reports/export/word', '导出 Word 报表', 'subject_id?', 'docx 文件', '401'),
        ('API-32', 'GET /api/reports/export/pdf', '导出 PDF 报表', 'subject_id?', 'pdf 文件', '401'),
        ('API-33', 'GET /api/backup/export', '导出 ZIP 数据包', '—', 'zip 文件', '401'),
        ('API-34', 'POST /api/backup/import', '导入数据包', 'multipart: file', '{stats}', '400 解析失败'),
    ]
    api_table = None
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == '编号' and '方法与路径' in table.rows[0].cells[1].text:
            api_table = table
            break
    if api_table is None:
        raise KeyError('API table not found')
    fill_table(api_table, ['编号', '方法与路径', '功能', '关键入参', '成功响应', '错误响应'], api_rows)
    set_repeat_table_header(api_table.rows[0])

    integrity_anchor = find_paragraph(doc, '数据完整性规则：')
    extension_tables = [
        ('5.8　review_plans（复习计划表）', [
            ('字段', '类型', '约束', '说明'),
            ('id', 'INTEGER', 'PK，自增', '计划 ID'),
            ('user_id', 'INTEGER', 'FK→users.id，NOT NULL，索引', '所属用户'),
            ('question_id', 'INTEGER', 'FK→questions.id，NOT NULL，索引', '关联错题'),
            ('plan_date', 'VARCHAR(10)', 'NOT NULL', '计划日期 YYYY-MM-DD'),
            ('status', 'VARCHAR(20)', '默认 pending', 'pending / done / skipped'),
            ('note', 'TEXT', '默认空', '复习备注'),
            ('created_at', 'DATETIME', '默认当前时间', '创建时间'),
        ]),
        ('5.9　ai_materials（知识库内容表）', [
            ('字段', '类型', '约束', '说明'),
            ('id', 'INTEGER', 'PK，自增', '内容 ID'),
            ('user_id', 'INTEGER', 'FK→users.id，NOT NULL，索引', '所属用户'),
            ('kp_id', 'INTEGER', 'FK→knowledge_points.id，NOT NULL', '知识点 ID'),
            ('type', 'VARCHAR(20)', 'NOT NULL', 'concept / variant / custom 等类型'),
            ('content', 'TEXT', 'NOT NULL', '知识库正文'),
            ('created_at', 'DATETIME', '默认当前时间', '创建时间'),
        ]),
        ('5.10　knowledge_traces（知识溯源记录表）', [
            ('字段', '类型', '约束', '说明'),
            ('id', 'INTEGER', 'PK，自增', '记录 ID'),
            ('user_id', 'INTEGER', 'FK→users.id，NOT NULL，索引', '所属用户'),
            ('question_id', 'INTEGER', 'FK→questions.id，NOT NULL，索引', '关联错题'),
            ('content', 'TEXT', 'NOT NULL', '结构化溯源结果'),
            ('created_at', 'DATETIME', '默认当前时间', '创建时间'),
        ]),
        ('5.11　mind_maps（知识溯源思维导图表）', [
            ('字段', '类型', '约束', '说明'),
            ('id', 'INTEGER', 'PK，自增', '导图 ID'),
            ('user_id', 'INTEGER', 'FK→users.id，NOT NULL，索引', '所属用户'),
            ('question_id', 'INTEGER', 'FK→questions.id，NOT NULL，索引', '关联错题'),
            ('data', 'TEXT', 'NOT NULL', '导图 JSON 数据'),
            ('created_at', 'DATETIME', '默认当前时间', '创建时间'),
            ('updated_at', 'DATETIME', '更新时自动刷新', '最后修改时间'),
        ]),
        ('5.12　kp_feedback（推荐反馈表）', [
            ('字段', '类型', '约束', '说明'),
            ('id', 'INTEGER', 'PK，自增', '反馈 ID'),
            ('user_id', 'INTEGER', 'FK→users.id，NOT NULL，索引', '所属用户'),
            ('question_id', 'INTEGER', 'FK→questions.id，可空', '来源错题'),
            ('kp_id', 'INTEGER', 'FK→knowledge_points.id，NOT NULL，索引', '知识点 ID'),
            ('action', 'VARCHAR(20)', 'NOT NULL', 'confirm / remove'),
            ('created_at', 'DATETIME', '默认当前时间', '创建时间'),
        ]),
        ('5.13　learning_resources（学习资源表）', [
            ('字段', '类型', '约束', '说明'),
            ('id', 'INTEGER', 'PK，自增', '资源 ID'),
            ('user_id', 'INTEGER', 'FK→users.id，NOT NULL，索引', '所属用户'),
            ('kp_id', 'INTEGER', 'FK→knowledge_points.id，NOT NULL，索引', '知识点 ID'),
            ('type', 'VARCHAR(20)', '默认 note', 'course / note / book / file 等类型'),
            ('title', 'VARCHAR(200)', '默认空', '资源标题'),
            ('url', 'VARCHAR(500)', '默认空', '资源链接'),
            ('page', 'VARCHAR(50)', '默认空', '教材页码'),
            ('content', 'TEXT', '默认空', '资源正文或说明'),
            ('attachment', 'VARCHAR(500)', '默认空', '附件相对路径'),
            ('created_at', 'DATETIME', '默认当前时间', '创建时间'),
        ]),
    ]
    for title, rows in extension_tables:
        add_before(doc, integrity_anchor, title, style='Heading 3')
        add_table_before(doc, integrity_anchor, rows[0], rows[1:])

    trc_anchor = find_paragraph(doc, '3.7　薄弱分析模块')
    trc_sections = [
        (
            'FR-TRC-06　知识溯源记录与思维导图',
            '功能描述：调用知识溯源接口后保存结构化溯源结果，并根据用户确认的知识点生成独立思维导图。',
            '接口：POST /api/questions/{qid}/trace、POST /api/questions/{qid}/trace/confirm、GET/POST /api/questions/{qid}/mindmap、GET /api/questions/{qid}/traces。',
            '验收标准：溯源结果可保存和再次读取；思维导图不写入知识树；导图节点可编辑并持久化。',
        ),
        (
            'FR-TRC-07　推荐反馈学习',
            '功能描述：记录用户确认或删除知识点候选的反馈，用于调整后续推荐权重；删除推荐不会影响其他用户。',
            '接口：POST /api/questions/{qid}/kp_feedback、POST /api/questions/{qid}/confirm_kp；反馈记录写入 kp_feedback 表。',
            '验收标准：确认和删除动作分别正确落库；再次推荐时反馈权重生效；不同用户反馈互不影响。',
        ),
    ]
    for title, description, spec, acceptance in trc_sections:
        add_before(doc, trc_anchor, title, style='Heading 4')
        add_before(doc, trc_anchor, description)
        add_before(doc, trc_anchor, spec)
        add_before(doc, trc_anchor, acceptance)
        add_table_before(
            doc,
            trc_anchor,
            ['优先级', 'P1', '实现状态', '已实现'],
            [],
            repeat_header=False,
        )

    routes = route_bindings()
    doc.add_page_break()
    doc.add_heading('附录 D　完整接口索引', level=2)
    doc.add_paragraph(
        f'以下清单由当前 app.py 自动核对生成，共 {len(routes)} 个路由路径。'
        '同一路径包含多个 HTTP 方法时，在方法列以逗号分隔。'
    )
    route_table = doc.add_table(rows=1, cols=3)
    route_table.style = 'Table Grid'
    fill_table(
        route_table,
        ['编号', '方法', '路径'],
        [(f'API-{index:02d}', ', '.join(methods), path) for index, (path, methods) in enumerate(routes, 1)],
    )
    set_repeat_table_header(route_table.rows[0])

    doc.save(SRS_PATH)
    print(SRS_PATH)


if __name__ == '__main__':
    main()
