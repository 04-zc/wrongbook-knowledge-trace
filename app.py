#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os
import io
import json
import re
import zipfile
import requests
import docx
from PyPDF2 import PdfReader
import fitz
from flask import Flask, request, jsonify, send_from_directory, session, send_file
from werkzeug.security import generate_password_hash, check_password_hash

import db

try:
    from paddleocr import PaddleOCR
    PADDLE_AVAILABLE = True
except Exception as e:
    print('PaddleOCR import failed:', e)
    PADDLE_AVAILABLE = False

_ocr = None
_ocr_cache = {}
def get_ocr():
    lang = db.get_setting('ocr_lang', 'ch') or 'ch'
    orientation = (db.get_setting('ocr_use_orientation', 'true') or 'true') == 'true'
    key = (lang, orientation)
    if key not in _ocr_cache and PADDLE_AVAILABLE:
        try:
            _ocr_cache[key] = PaddleOCR(use_textline_orientation=orientation, lang=lang)
        except TypeError:
            # PaddleOCR 2.x 使用旧参数
            _ocr_cache[key] = PaddleOCR(use_angle_cls=orientation, lang=lang, show_log=False)
    return _ocr_cache.get(key)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, 'static')
UPLOAD_DIR = os.path.join(BASE_DIR, 'uploads')
os.makedirs(UPLOAD_DIR, exist_ok=True)

app = Flask(__name__, static_folder=STATIC_DIR, static_url_path='/static')
app.config['JSON_AS_ASCII'] = False

SECRET_FILE = os.path.join(BASE_DIR, 'secret.key')
if os.path.exists(SECRET_FILE):
    with open(SECRET_FILE, 'rb') as f:
        app.secret_key = f.read()
else:
    app.secret_key = os.urandom(32)
    with open(SECRET_FILE, 'wb') as f:
        f.write(app.secret_key)

@app.before_request
def require_login():
    # 静态资源、首页与登录接口不拦截
    if request.path.startswith('/static/') or request.path.startswith('/uploads/') or request.path == '/':
        return None
    if request.path.startswith('/api/auth/'):
        return None
    if 'user_id' not in session:
        return jsonify({'error': '未登录', 'auth_required': True}), 401
    db.set_current_user(session.get('user_id'))
    return None

@app.after_request
def add_cors(resp):
    resp.headers['Access-Control-Allow-Origin'] = '*'
    resp.headers['Access-Control-Allow-Methods'] = 'GET, POST, PUT, DELETE, OPTIONS'
    resp.headers['Access-Control-Allow-Headers'] = 'Content-Type'
    # 开发阶段禁用静态文件缓存，避免前端改了不生效
    if request.path.startswith('/static/'):
        resp.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        resp.headers['Pragma'] = 'no-cache'
    return resp

# ---------- helpers ----------
def extract_text_for_ai(q):
    parts = []
    for k in ['title_text', 'analysis', 'correct_answer']:
        if q.get(k):
            parts.append(q[k])
    return '\n'.join(parts)

def local_recommend(text, subject_id):
    kps = db.list_kps(subject_id)
    matches = []
    for kp in kps:
        if kp['name'] in text:
            matches.append({
                'id': kp['id'],
                'path': kp['name'],
                'name': kp['name'],
                'confidence': 0.85,
                'is_auto': 1
            })
    keyword_map = {
        '诱导公式': '三角函数',
        '三角函数': '三角函数',
        '圆的性质': '平面几何',
        '平面几何': '平面几何',
        '现在完成时': '时态',
        '过去式': '时态',
        '函数': '函数',
        '导数': '函数',
        '积分': '函数'
    }
    for kw, parent in keyword_map.items():
        if kw in text:
            for kp in kps:
                if kp['name'] == parent and not any(m['id'] == kp['id'] for m in matches):
                    matches.append({
                        'id': kp['id'],
                        'path': kp['name'],
                        'name': kp['name'],
                        'confidence': 0.75,
                        'is_auto': 1
                    })
    return matches[:3]

def call_llm(prompt, json_mode=False):
    api_key = db.get_setting('llm_api_key', '')
    model = db.get_setting('llm_model', 'deepseek-chat')
    if not api_key:
        raise RuntimeError('未配置大模型 API Key，请点击右上角「大模型配置」填写')
    if 'deepseek' in model:
        url = 'https://api.deepseek.com/chat/completions'
        payload = {
            'model': model or 'deepseek-chat',
            'messages': [{'role': 'user', 'content': prompt}],
            'temperature': 0.4
        }
    else:
        url = 'https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions'
        payload = {
            'model': model or 'qwen-turbo',
            'messages': [{'role': 'user', 'content': prompt}],
            'temperature': 0.4
        }
    if json_mode:
        payload['response_format'] = {'type': 'json_object'}
    resp = requests.post(url, json=payload, headers={
        'Content-Type': 'application/json',
        'Authorization': f'Bearer {api_key}'
    }, timeout=60)
    resp.raise_for_status()
    return resp.json()['choices'][0]['message']['content']

def classify_recommendations(recs):
    boost = db.feedback_boost_map()
    try:
        auto_th = float(db.get_setting('kp_auto_threshold', '0.8') or 0.8)
    except (TypeError, ValueError):
        auto_th = 0.8
    try:
        pending_th = float(db.get_setting('kp_pending_threshold', '0.5') or 0.5)
    except (TypeError, ValueError):
        pending_th = 0.5
    result = []
    for rec in recs:
        item = dict(rec)
        kp_id = item.get('id')
        confidence = float(item.get('confidence') or 0) + boost.get(kp_id, 0)
        confidence = min(1.0, max(0.0, confidence))
        item['confidence'] = round(confidence, 2)
        item['feedback_boost'] = round(boost.get(kp_id, 0), 2)
        if confidence >= auto_th:
            item['decision'] = 'auto'
        elif confidence >= pending_th:
            item['decision'] = 'pending'
        else:
            item['decision'] = 'ignored'
        if item['decision'] != 'ignored':
            result.append(item)
    return result

def ai_recommend(text, subject_id):
    if not db.get_setting('llm_api_key', ''):
        return classify_recommendations(local_recommend(text, subject_id))

    prompt = f'''请分析以下错题，判断它考查的知识点，按"学科 > 一级知识点 > 二级知识点 > 三级知识点"的格式返回。
如果涉及多个知识点，请列出最多 3 个。
同时给出每个知识点的置信度（0~1）。
path 中每一级都只写精简名称（12 个字以内），不要写解释。
只返回 JSON，格式如下：
{{"knowledge_points":[{{"path":"数学 > 函数 > 三角函数 > 诱导公式","confidence":0.92}}]}}

题目：{text}
'''
    try:
        parsed = json.loads(call_llm(prompt, json_mode=True))
        result = []
        kps = db.list_kps(subject_id)
        for item in parsed.get('knowledge_points', []):
            path = item.get('path', '')
            name = path.split('>')[-1].strip() if '>' in path else path.strip()
            matched = None
            for kp in kps:
                if kp['name'] == name:
                    matched = kp
                    break
            if not matched and name:
                matched = db.add_kp(subject_id, name, description=f'来自AI推荐: {path}')
                kps.append(matched)
            if matched:
                result.append({
                    'id': matched['id'],
                    'path': path,
                    'name': matched['name'],
                    'confidence': item.get('confidence', 0.8),
                    'is_auto': 1
                })
        return classify_recommendations(result)
    except Exception as e:
        print('AI recommend error:', e)
        return classify_recommendations(local_recommend(text, subject_id))

QUESTION_MARK = re.compile(
    r'(?m)^\s*(?:\d{1,3}\s*[.、．)）]|[（(]\d{1,3}[）)]|第\s*\d{1,3}\s*题)'
)

MATH_STYLE = '''

【排版要求】请输出人类可以直接阅读的纯文本，不要输出 LaTeX 源码、Markdown 表格或代码块：
1. 保留自然换行和分步过程。
2. 分数写成 a ÷ b 或 (a)/(b)，不要写 \\frac。
3. 乘号用 ×，点乘用 ·，除号用 ÷。
4. 平方、立方用 x²、x³，下标用 x₁、v₀。
5. 希腊字母直接用 α β θ π λ μ Δ ω ρ σ φ ε 等。
6. 物理量、向量用 F⃗、v⃗ 这种可读写法，单位写成 m/s²、N·m。
7. 极限、积分、求和写成 lim(x→0)、∫、Σ 并保留上下限文字说明。
8. 矩阵按行换行并用 [ ] 包裹，例如：
[ 1  2 ]
[ 3  4 ]
'''

def split_questions(text):
    text = (text or '').replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'\n{3,}', '\n\n', text).strip()
    if not text:
        return []
    matches = list(QUESTION_MARK.finditer(text))
    parts = []
    if matches:
        starts = [m.start() for m in matches] + [len(text)]
        for i in range(len(matches)):
            seg = text[starts[i]:starts[i + 1]].strip()
            seg = QUESTION_MARK.sub('', seg, count=1).strip()
            if len(seg) >= 4:
                parts.append(seg)
    if not parts:
        blocks = [b.strip() for b in re.split(r'\n\s*\n', text) if len(b.strip()) >= 6]
        parts = blocks if blocks else ([text] if len(text) >= 6 else [])
    return [{'title_text': p} for p in parts[:80]]

def extract_docx_text(path):
    document = docx.Document(path)
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(' | '.join(cells))
    return '\n'.join(lines)

def extract_pdf_text(path):
    doc = fitz.open(path)
    page_count = len(doc)
    text = '\n'.join(page.get_text() for page in doc)
    ocr_used = False
    # 文本过少说明是扫描版，转图片走 OCR
    if len(text.strip()) < max(10, page_count * 10):
        ocr_used = True
        texts = []
        for idx, page in enumerate(doc):
            if idx >= 10:
                break
            pix = page.get_pixmap(dpi=200)
            img_path = os.path.join(UPLOAD_DIR, f'pdf_page_{os.urandom(3).hex()}.png')
            pix.save(img_path)
            try:
                texts.append(ocr_image(img_path))
            finally:
                try:
                    os.remove(img_path)
                except OSError:
                    pass
        text = '\n'.join(texts)
    doc.close()
    return text, ocr_used

# ---------- routes ----------
@app.route('/api/auth/me', methods=['GET'])
def auth_me():
    uid = session.get('user_id')
    if not uid:
        return jsonify({'user': None}), 401
    u = db.get_user(uid)
    if not u:
        session.clear()
        return jsonify({'user': None}), 401
    return jsonify({'user': {'username': u['username']}})

@app.route('/api/auth/register', methods=['POST'])
def auth_register():
    body = request.get_json(force=True, silent=True) or {}
    username = (body.get('username') or '').strip()
    password = body.get('password') or ''
    if not re.fullmatch(r'[A-Za-z0-9]{3,20}', username):
        return jsonify({'error': '账号只能是 3-20 位英文字母或数字'}), 400
    if not re.fullmatch(r'[A-Za-z0-9]{6,20}', password):
        return jsonify({'error': '密码只能是 6-20 位英文字母或数字'}), 400
    if db.get_user_by_username(username):
        return jsonify({'error': '该账号已存在，请直接登录或换一个账号'}), 400
    was_empty = db.count_users() == 0
    u = db.create_user(username, generate_password_hash(password))
    if was_empty:
        # 首个账号接管旧版全局数据，之后的账号各自从空库开始
        db.claim_legacy_data(u['id'])
    session['user_id'] = u['id']
    session['username'] = u['username']
    return jsonify({'user': {'username': u['username']}})

@app.route('/api/auth/login', methods=['POST'])
def auth_login():
    body = request.get_json(force=True, silent=True) or {}
    username = (body.get('username') or '').strip()
    password = body.get('password') or ''
    u = db.get_user_by_username(username)
    if not u:
        return jsonify({'error': '账号不存在'}), 400
    stored_hash = db.get_user_password(username)
    if not check_password_hash(stored_hash, password):
        return jsonify({'error': '密码不正确'}), 400
    session['user_id'] = u['id']
    session['username'] = u['username']
    db.set_current_user(u['id'])
    return jsonify({'user': {'username': u['username']}})

@app.route('/api/auth/logout', methods=['POST'])
def auth_logout():
    session.clear()
    return jsonify({'ok': True})

@app.route('/api/auth/change_password', methods=['POST'])
def auth_change_password():
    uid = session.get('user_id')
    body = request.get_json(force=True, silent=True) or {}
    old = body.get('old_password') or ''
    new = body.get('new_password') or ''
    if not re.fullmatch(r'[A-Za-z0-9]{6,20}', new):
        return jsonify({'error': '新密码只能是 6-20 位英文字母或数字'}), 400
    stored = db.get_user_password_by_id(uid)
    if not stored or not check_password_hash(stored, old):
        return jsonify({'error': '原密码不正确'}), 400
    db.update_user_password(uid, generate_password_hash(new))
    return jsonify({'ok': True})

@app.route('/api/auth/reset_password', methods=['POST'])
def auth_reset_password():
    body = request.get_json(force=True, silent=True) or {}
    username = (body.get('username') or '').strip()
    new = body.get('new_password') or ''
    user = db.get_user_by_username(username)
    if not user:
        return jsonify({'error': '账号不存在'}), 404
    if not re.fullmatch(r'[A-Za-z0-9]{6,20}', new):
        return jsonify({'error': '新密码只能是 6-20 位英文字母或数字'}), 400
    db.update_user_password(user['id'], generate_password_hash(new))
    return jsonify({'ok': True})

@app.route('/api/auth/rename', methods=['POST'])
def auth_rename():
    uid = session.get('user_id')
    body = request.get_json(force=True, silent=True) or {}
    new_username = (body.get('new_username') or '').strip()
    password = body.get('password') or ''
    if not re.fullmatch(r'[A-Za-z0-9]{3,20}', new_username):
        return jsonify({'error': '账号只能是 3-20 位英文字母或数字'}), 400
    stored = db.get_user_password_by_id(uid)
    if not stored or not check_password_hash(stored, password):
        return jsonify({'error': '密码不正确'}), 400
    if not db.rename_user(uid, new_username):
        return jsonify({'error': '该账号已存在'}), 400
    session['username'] = new_username
    return jsonify({'ok': True, 'username': new_username})

@app.route('/api/auth/account', methods=['DELETE'])
def auth_delete_account():
    uid = session.get('user_id')
    body = request.get_json(force=True, silent=True) or {}
    password = body.get('password') or ''
    stored = db.get_user_password_by_id(uid)
    if not stored or not check_password_hash(stored, password):
        return jsonify({'error': '密码不正确'}), 400
    db.delete_user_account(uid)
    session.clear()
    return jsonify({'ok': True})

@app.route('/')
def index():
    return send_from_directory(STATIC_DIR, 'index.html')

@app.route('/uploads/<path:filename>')
def uploaded_file(filename):
    return send_from_directory(UPLOAD_DIR, filename)

@app.route('/api/import/document', methods=['POST'])
def api_import_document():
    if 'file' not in request.files:
        return jsonify({'error': '请选择 PDF 或 Word 文件'}), 400
    file = request.files['file']
    filename = file.filename or ''
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ('.pdf', '.docx', '.doc'):
        return jsonify({'error': '只支持 PDF、DOC、DOCX 文件'}), 400
    save_name = f"import_{os.urandom(4).hex()}{ext}"
    save_path = os.path.join(UPLOAD_DIR, save_name)
    file.save(save_path)
    try:
        if ext == '.pdf':
            text, ocr_used = extract_pdf_text(save_path)
            mode = 'OCR 识别' if ocr_used else '文本提取'
        else:
            text = extract_docx_text(save_path)
            mode = '文本提取'
        questions = split_questions(text)
        return jsonify({
            'filename': filename,
            'mode': mode,
            'count': len(questions),
            'questions': questions
        })
    except Exception as e:
        return jsonify({'error': f'文件解析失败：{e}'}), 500

@app.route('/api/import/save', methods=['POST'])
def api_import_save():
    body = request.get_json(force=True, silent=True) or {}
    subject_id = body.get('subject_id')
    questions = body.get('questions') or []
    if not subject_id:
        return jsonify({'error': '请选择学科'}), 400
    if not questions:
        return jsonify({'error': '没有可导入的题目'}), 400
    saved = 0
    for item in questions:
        title = (item.get('title_text') or '').strip()
        if not title:
            continue
        db.add_question({
            'subject_id': subject_id,
            'title_text': title,
            'source': 'import',
            'status': 0,
            'knowledge_points': []
        })
        saved += 1
    return jsonify({'saved': saved})

# ---------- ai materials ----------
@app.route('/api/ai/generate', methods=['POST'])
def api_ai_generate():
    body = request.get_json(force=True, silent=True) or {}
    kp_id = body.get('kp_id')
    mode = body.get('mode')
    count = int(body.get('count') or 2)
    kp = db.get_kp(kp_id) if kp_id else None
    if mode != 'explain' and not kp:
        return jsonify({'error': '知识点不存在'}), 404
    if mode == 'concept':
        prompt = f'''请针对知识点「{kp['name']}」生成概念回顾：
1. 用 200 字左右概括核心概念和公式。
2. 列出 3 个易错点。
直接输出内容，不要多余说明。''' + MATH_STYLE
    elif mode == 'variant':
        count = min(max(count, 1), 5)
        prompt = f'''请针对知识点「{kp['name']}」出 {count} 道变式练习题。
只返回 JSON，格式：
{{"questions":[{{"question":"题目","answer":"答案","analysis":"解析"}}]}}
题目和公式要写成适合人类阅读的形式：
1. 保留自然换行，不要挤成一行。
2. 除法用 "÷" 或文字表达，不要用 "/"；乘法用 "×"。
3. 不要使用 Markdown 表格、代码块或转义符号。
4. 不要编造试卷地市和年份。''' + MATH_STYLE
    elif mode == 'explain':
        q = db.get_question(body.get('question_id'))
        if not q:
            return jsonify({'error': '错题不存在'}), 404
        prompt = f'''请讲解下面这道错题，输出：
1. 这道题考查的知识点
2. 正确的解题步骤
3. 学生容易错在哪里
4. 同类题的解题提醒

题目：{q['title_text']}
正确答案：{q.get('correct_answer') or '未填写'}
学生答案：{q.get('user_answer') or '未填写'}
已有解析：{q.get('analysis') or '无'}''' + MATH_STYLE
    elif mode == 'typical':
        prompt = f'''请针对知识点「{kp['name']}」给出 2 道典型例题。
只返回 JSON，格式：
{{"questions":[{{"question":"题目","answer":"答案","analysis":"解析与易错点"}}]}}
题目和公式要写成适合人类阅读的形式：
1. 保留自然换行，不要挤成一行。
2. 除法用 "÷" 或文字表达，不要用 "/"；乘法用 "×"。
3. 不要使用 Markdown 表格、代码块或转义符号。
4. 不要编造试卷地市和年份。''' + MATH_STYLE
    elif mode == 'summary':
        prompt = f'''请把知识点「{kp['name']}」整理成一份精简、可直接用于复习的知识卡片：
1. 核心概念，不超过 100 字。
2. 必背公式或定律。
3. 最常见题型的解题方法，写成步骤。
4. 2 到 3 个易错点。
要求去重、去废话，只保留最有用的内容，直接输出。''' + MATH_STYLE
    elif mode == 'resources':
        prompt = f'''请为知识点「{kp['name']}」推荐学习资源，只返回 JSON：
{{"resources":[
  {{"type":"course","title":"B站学习视频","url":"https://search.bilibili.com/all?keyword=关键词","content":"适合什么时候看、重点看什么"}},
  {{"type":"note","title":"复习笔记","content":"精简的复习笔记内容"}}
]}}
要求：
1. 至少 1 条 B站网课资源，url 只能使用 B站搜索链接，不要编造 BV 号。
2. 至少 1 条复习笔记，笔记控制在 200 字以内。
3. 不要推荐教材页码。
4. 不要使用 Markdown 或代码块。''' + MATH_STYLE
    else:
        return jsonify({'error': '不支持的类型'}), 400
    try:
        content = call_llm(prompt, json_mode=(mode in ('typical', 'variant', 'resources')))
        if mode in ('typical', 'variant'):
            try:
                questions = json.loads(content).get('questions', [])
            except Exception:
                questions = [{'question': content, 'answer': '', 'analysis': ''}]
            return jsonify({'mode': mode, 'questions': questions})
        if mode == 'resources':
            try:
                resources = json.loads(content).get('resources', [])
            except Exception:
                resources = []
            return jsonify({'mode': mode, 'resources': resources})
        return jsonify({'mode': mode, 'content': content})
    except RuntimeError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        return jsonify({'error': f'生成失败：{e}'}), 500

@app.route('/api/ai/save', methods=['POST'])
def api_ai_save():
    body = request.get_json(force=True, silent=True) or {}
    if not body.get('kp_id') or not body.get('content'):
        return jsonify({'error': '内容为空'}), 400
    m = db.add_ai_material(body['kp_id'], body.get('mode', 'concept'), body['content'])
    return jsonify({'ok': True, 'material': m})

@app.route('/api/ai/materials', methods=['GET'])
def api_ai_materials():
    return jsonify(db.list_ai_materials(request.args.get('kp_id', type=int)))

@app.route('/api/ai/materials', methods=['POST'])
def api_create_ai_material():
    body = request.get_json(force=True, silent=True) or {}
    if not body.get('kp_id') or not (body.get('content') or '').strip():
        return jsonify({'error': '知识点和内容不能为空'}), 400
    result = db.add_ai_material(
        body['kp_id'],
        body.get('mode') or body.get('type') or 'custom',
        body['content']
    )
    return jsonify({'ok': True, 'material': result})

@app.route('/api/ai/materials/<int:mid>', methods=['POST'])
def api_update_ai_material(mid):
    body = request.get_json(force=True, silent=True) or {}
    result = db.update_ai_material(mid, body.get('content', ''))
    if not result:
        return jsonify({'error': '内容不存在'}), 404
    return jsonify({'ok': True, 'material': result})

@app.route('/api/ai/materials/<int:mid>', methods=['DELETE'])
def api_delete_ai_material(mid):
    db.delete_ai_material(mid)
    return jsonify({'ok': True})

@app.route('/api/kp/typical', methods=['GET'])
def api_kp_typical():
    kp_id = request.args.get('kp_id', type=int)
    if not kp_id:
        return jsonify({'error': 'kp_id required'}), 400
    return jsonify(db.typical_examples(kp_id, limit=5))

@app.route('/api/knowledge_graph', methods=['GET'])
def api_knowledge_graph():
    subject_id = request.args.get('subject_id', type=int)
    if not subject_id:
        return jsonify({'error': 'subject_id required'}), 400
    subjects = [s for s in db.list_subjects() if s['id'] == subject_id]
    if not subjects:
        return jsonify({'error': '学科不存在'}), 404
    return jsonify({'subject': subjects[0], 'nodes': db.kp_tree(subject_id)})

@app.route('/api/analytics/trend', methods=['GET'])
def api_trend():
    return jsonify(db.trend_data(
        subject_id=request.args.get('subject_id', type=int),
        days=request.args.get('days', 30, type=int)
    ))

@app.route('/api/analytics/priority', methods=['GET'])
def api_priority():
    subject_id = request.args.get('subject_id', type=int)
    rows = [kp for kp in db.weak_analysis(subject_id=subject_id, top_n=10) if kp.get('question_count')]
    result = []
    for kp in rows:
        score = kp.get('weak_index', 0)
        if score >= 7:
            level, action = '高', '优先复习：先回顾概念，再做变式练习'
        elif score >= 4:
            level, action = '中', '安排近期复习：重做错题并总结方法'
        else:
            level, action = '低', '保持巩固：定期快速回顾'
        result.append({**kp, 'priority': level, 'suggestion': action})
    return jsonify(result)

@app.route('/api/reports/summary', methods=['GET'])
def api_report_summary():
    subject_id = request.args.get('subject_id', type=int)
    questions = db.list_questions(subject_id=subject_id)
    status_counts = {0: 0, 1: 0, 2: 0, 3: 0}
    for q in questions:
        status_counts[q.get('status', 0)] = status_counts.get(q.get('status', 0), 0) + 1
    trend = db.trend_data(subject_id=subject_id, days=30)
    weak = [item for item in db.weak_analysis(subject_id=subject_id, top_n=10) if item.get('question_count')]
    return jsonify({
        'subject_id': subject_id,
        'totals': {
            'subjects': len([s for s in db.list_subjects() if not subject_id or s['id'] == subject_id]),
            'knowledge_points': len(db.list_kps(subject_id)),
            'questions': len(questions),
            'status_counts': status_counts,
            'added_30d': sum(item['added'] for item in trend),
            'wrong_reviews_30d': sum(item['wrong_reviews'] for item in trend),
            'review_plans': db.review_plan_counts(),
            'reviews': db.review_stats(subject_id)
        },
        'weak_points': weak,
        'trend': trend
    })

# subjects
@app.route('/api/subjects', methods=['GET'])
def api_subjects():
    return jsonify(db.list_subjects())

@app.route('/api/subjects', methods=['POST'])
def api_add_subject():
    body = request.get_json(force=True, silent=True) or {}
    return jsonify(db.add_subject(body.get('name', ''), body.get('icon', '')))

@app.route('/api/subjects/<int:sid>', methods=['DELETE'])
def api_del_subject(sid):
    db.delete_subject(sid)
    return jsonify({'ok': True})

# knowledge points
@app.route('/api/knowledge_points', methods=['GET'])
def api_kps():
    subject_id = request.args.get('subject_id', type=int)
    return jsonify(db.list_kps(subject_id))

@app.route('/api/knowledge_points', methods=['POST'])
def api_add_kp():
    body = request.get_json(force=True, silent=True) or {}
    return jsonify(db.add_kp(
        body.get('subject_id'),
        body.get('name'),
        body.get('parent_id'),
        body.get('level', 0),
        body.get('description', '')
    ))

@app.route('/api/knowledge_points/<int:kid>', methods=['DELETE'])
def api_del_kp(kid):
    db.delete_kp(kid)
    return jsonify({'ok': True})

@app.route('/api/knowledge_points/<int:kid>/description', methods=['POST'])
def api_update_kp_description(kid):
    body = request.get_json(force=True, silent=True) or {}
    db.update_kp_description(kid, body.get('description', ''))
    return jsonify({'ok': True})

@app.route('/api/knowledge_points/<int:kid>/resources', methods=['GET'])
def api_list_resources(kid):
    return jsonify(db.list_learning_resources(kid))

@app.route('/api/knowledge_points/<int:kid>/resources', methods=['POST'])
def api_add_resource(kid):
    body = request.get_json(force=True, silent=True) or {}
    return jsonify(db.add_learning_resource(kid, body))

@app.route('/api/resources/<int:rid>', methods=['POST'])
def api_update_resource(rid):
    body = request.get_json(force=True, silent=True) or {}
    result = db.update_learning_resource(rid, body)
    if not result:
        return jsonify({'error': '资源不存在'}), 404
    return jsonify({'ok': True, 'resource': result})

@app.route('/api/resources/<int:rid>', methods=['DELETE'])
def api_delete_resource(rid):
    db.delete_learning_resource(rid)
    return jsonify({'ok': True})

@app.route('/api/upload_attachment', methods=['POST'])
def api_upload_attachment():
    if 'file' not in request.files:
        return jsonify({'error': '没有文件'}), 400
    file = request.files['file']
    if not file or not file.filename:
        return jsonify({'error': '没有文件'}), 400
    ext = os.path.splitext(file.filename)[1]
    save_name = f"attach_{os.urandom(5).hex()}{ext}"
    file.save(os.path.join(UPLOAD_DIR, save_name))
    return jsonify({'url': f'/uploads/{save_name}', 'filename': file.filename})

# questions
@app.route('/api/questions', methods=['GET'])
def api_questions():
    return jsonify(db.list_questions(
        subject_id=request.args.get('subject_id', type=int),
        status=request.args.get('status', type=int),
        kp_id=request.args.get('kp_id', type=int)
    ))

@app.route('/api/questions', methods=['POST'])
def api_add_question():
    body = request.get_json(force=True, silent=True) or {}
    return jsonify(db.add_question(body))

@app.route('/api/questions/<int:qid>', methods=['GET'])
def api_get_question(qid):
    q = db.get_question(qid)
    if not q:
        return jsonify({'error': 'not found'}), 404
    return jsonify(q)

@app.route('/api/questions/<int:qid>', methods=['PUT'])
def api_update_question(qid):
    body = request.get_json(force=True, silent=True) or {}
    q = db.update_question(qid, body)
    if not q:
        return jsonify({'error': 'not found'}), 404
    return jsonify(q)

@app.route('/api/questions/<int:qid>', methods=['DELETE'])
def api_del_question(qid):
    db.delete_question(qid)
    return jsonify({'ok': True})

@app.route('/api/questions/trash', methods=['GET'])
def api_trash_questions():
    return jsonify(db.list_trash_questions())

@app.route('/api/questions/batch', methods=['POST'])
def api_batch_questions():
    body = request.get_json(force=True, silent=True) or {}
    action = body.get('action')
    ids = [int(i) for i in (body.get('ids') or [])]
    if not ids:
        return jsonify({'error': '请先选择错题'}), 400
    if action == 'delete':
        for qid in ids:
            db.delete_question(qid)
        return jsonify({'ok': True, 'count': len(ids)})
    if action == 'status':
        status = int(body.get('status', 0))
        for qid in ids:
            db.update_question(qid, {'status': status})
        return jsonify({'ok': True, 'count': len(ids)})
    if action == 'bind_kp':
        kp_ids = [int(i) for i in (body.get('kp_ids') or [])]
        if not kp_ids:
            return jsonify({'error': '请选择知识点'}), 400
        kp_map = {kp['id']: kp for kp in db.list_kps()}
        linked = 0
        for qid in ids:
            q = db.get_question(qid)
            if not q:
                continue
            points = list(q.get('knowledge_points') or [])
            have = {p['id'] for p in points}
            for kid in kp_ids:
                if kid in have or kid not in kp_map:
                    continue
                kp = kp_map[kid]
                points.append({
                    'id': kp['id'],
                    'name': kp['name'],
                    'confidence': 1,
                    'is_auto': 0
                })
            db.update_question(qid, {'knowledge_points': points})
            linked += 1
        return jsonify({'ok': True, 'count': linked})
    return jsonify({'error': '不支持的操作'}), 400

@app.route('/api/questions/<int:qid>/restore', methods=['POST'])
def api_restore_question(qid):
    db.restore_question(qid)
    return jsonify({'ok': True})

@app.route('/api/questions/<int:qid>/purge', methods=['DELETE'])
def api_purge_question(qid):
    db.purge_question(qid)
    return jsonify({'ok': True})

@app.route('/api/questions/<int:qid>/recommend_kp', methods=['POST'])
def api_recommend_kp(qid):
    q = db.get_question(qid)
    if not q:
        return jsonify({'error': 'not found'}), 404
    text = extract_text_for_ai(q)
    result = ai_recommend(text, q['subject_id'])
    auto = [r for r in result if r.get('decision') == 'auto']
    pending = [r for r in result if r.get('decision') == 'pending']
    existing = list(q.get('knowledge_points') or [])
    have = {kp['id'] for kp in existing}
    for rec in auto:
        if rec['id'] in have:
            continue
        existing.append({
            'id': rec['id'],
            'name': rec['name'],
            'confidence': rec['confidence'],
            'is_auto': 1
        })
        have.add(rec['id'])
    if auto:
        db.update_question(qid, {'knowledge_points': existing})
    return jsonify({
        'recommendations': result,
        'auto_bound': auto,
        'pending': pending
    })

@app.route('/api/questions/<int:qid>/confirm_kp', methods=['POST'])
def api_confirm_kp(qid):
    q = db.get_question(qid)
    if not q:
        return jsonify({'error': 'not found'}), 404
    body = request.get_json(force=True, silent=True) or {}
    kp_ids = [int(i) for i in (body.get('kp_ids') or [])]
    points = list(q.get('knowledge_points') or [])
    have = {kp['id'] for kp in points}
    kp_map = {kp['id']: kp for kp in db.list_kps(q['subject_id'])}
    added = []
    for kid in kp_ids:
        if kid in have or kid not in kp_map:
            continue
        kp = kp_map[kid]
        points.append({
            'id': kp['id'],
            'name': kp['name'],
            'confidence': 1,
            'is_auto': 0
        })
        have.add(kid)
        added.append(kp)
        db.add_kp_feedback(qid, kid, 'confirm')
    if added:
        db.update_question(qid, {'knowledge_points': points})
    return jsonify({'ok': True, 'added': added})

@app.route('/api/questions/<int:qid>/kp_feedback', methods=['POST'])
def api_kp_feedback(qid):
    body = request.get_json(force=True, silent=True) or {}
    kp_id = body.get('kp_id')
    action = body.get('action')
    if not kp_id or action not in ('confirm', 'remove'):
        return jsonify({'error': '参数错误'}), 400
    db.add_kp_feedback(qid, int(kp_id), action)
    return jsonify({'ok': True})

@app.route('/api/questions/<int:qid>/trace', methods=['POST'])
def api_question_trace(qid):
    q = db.get_question(qid)
    if not q:
        return jsonify({'error': '错题不存在'}), 404
    if not db.get_setting('llm_api_key', ''):
        return jsonify({'error': '未配置大模型 API Key，请先在右上角「大模型配置」填写'}), 400
    prompt = f'''请对下面这道错题做知识溯源，只返回 JSON，格式：
{{
  "symptom": "错误表象，简短描述学生错在哪里",
  "direct_points": [{{"name": "直接考查的知识点名称", "confidence": 0.9, "reason": "详细说明，包含判断依据和内容"}}],
  "prerequisite_points": [{{"name": "前置知识点名称", "confidence": 0.8, "reason": "详细说明，包含判断依据和内容"}}],
  "root_cause": "最可能的根本原因",
  "path": ["错误表象", "直接知识点", "前置知识点", "根本原因"]
}}
直接知识点最多 3 个，前置知识点最多 3 个。
name 必须是精简的知识点名称，控制在 12 个字以内，只写名称本身，不要写解释。
详细解释、公式、判断依据全部写在 reason 里，后续会展示在知识点详情中。
题目：{q['title_text']}
正确答案：{q.get('correct_answer') or '未填写'}
学生答案：{q.get('user_answer') or '未填写'}
解析：{q.get('analysis') or '无'}
错因：{q.get('mistake_reason') or '无'}''' + MATH_STYLE
    try:
        parsed = json.loads(call_llm(prompt, json_mode=True))
    except Exception as e:
        return jsonify({'error': f'溯源失败：{e}'}), 500
    trace = db.add_trace(qid, json.dumps(parsed, ensure_ascii=False))
    existing = {kp['name'] for kp in db.list_kps(q['subject_id'])}
    for group in ('direct_points', 'prerequisite_points'):
        for item in parsed.get(group, []):
            item['exists'] = item.get('name') in existing
    parsed['trace_id'] = trace['id']
    parsed['question_id'] = qid
    return jsonify(parsed)

@app.route('/api/questions/<int:qid>/trace/confirm', methods=['POST'])
def api_question_trace_confirm(qid):
    q = db.get_question(qid)
    if not q:
        return jsonify({'error': '错题不存在'}), 404
    body = request.get_json(force=True, silent=True) or {}
    points = body.get('points') or []
    nodes = [{
        'id': 'question',
        'type': 'question',
        'label': '这道错题',
        'detail': q['title_text']
    }]
    edges = []
    if body.get('symptom'):
        nodes.append({'id': 'symptom', 'type': 'symptom', 'label': '错误表象', 'detail': body['symptom']})
        edges.append({'from': 'question', 'to': 'symptom'})
    last_direct_ids = []
    for idx, item in enumerate(points):
        name = (item.get('name') or '').strip()
        if not name:
            continue
        node_id = f"point-{idx}"
        is_prerequisite = item.get('group') == 'prerequisite_points'
        nodes.append({
            'id': node_id,
            'type': 'prerequisite' if is_prerequisite else 'direct',
            'label': name,
            'detail': item.get('reason') or '',
            'confidence': item.get('confidence', 0)
        })
        edges.append({'from': 'symptom' if body.get('symptom') else 'question', 'to': node_id})
        if not is_prerequisite:
            last_direct_ids.append(node_id)
    if body.get('root_cause'):
        nodes.append({'id': 'root-cause', 'type': 'root', 'label': '根本原因', 'detail': body['root_cause']})
        for node_id in (last_direct_ids or ['symptom' if body.get('symptom') else 'question']):
            edges.append({'from': node_id, 'to': 'root-cause'})
    map_data = {
        'question_id': qid,
        'question_text': q['title_text'],
        'nodes': nodes,
        'edges': edges
    }
    saved = db.save_mind_map(qid, json.dumps(map_data, ensure_ascii=False))
    return jsonify({'ok': True, 'mindmap': map_data, 'saved': saved})

@app.route('/api/questions/<int:qid>/mindmap', methods=['GET'])
def api_get_mindmap(qid):
    saved = db.get_mind_map(qid)
    if not saved:
        return jsonify({'mindmap': None})
    return jsonify({'mindmap': json.loads(saved['data']), 'updated_at': saved['updated_at']})

@app.route('/api/questions/<int:qid>/mindmap', methods=['POST'])
def api_save_mindmap(qid):
    body = request.get_json(force=True, silent=True) or {}
    map_data = body.get('mindmap')
    if not map_data:
        return jsonify({'error': '导图内容为空'}), 400
    saved = db.save_mind_map(qid, json.dumps(map_data, ensure_ascii=False))
    return jsonify({'ok': True, 'saved': saved})

@app.route('/api/questions/<int:qid>/traces', methods=['GET'])
def api_question_traces(qid):
    return jsonify(db.list_traces(qid))

@app.route('/api/questions/<int:qid>/review', methods=['POST'])
def api_review(qid):
    body = request.get_json(force=True, silent=True) or {}
    db.add_review(qid, body.get('result', 0), body.get('note', ''))
    return jsonify({'ok': True})

# weak analysis
@app.route('/api/weak_points', methods=['GET'])
def api_weak():
    return jsonify(db.weak_analysis(
        subject_id=request.args.get('subject_id', type=int),
        top_n=request.args.get('top_n', 10, type=int)
    ))

@app.route('/api/review/today', methods=['GET'])
def api_review_today():
    subject_id = request.args.get('subject_id', type=int)
    all_kps = db.kp_tree(subject_id) if subject_id else []
    recommendations = [
        kp for kp in db.weak_analysis(subject_id=subject_id, top_n=10)
        if (kp.get('question_count') or 0) > 0
    ][:5]
    if recommendations:
        names = '、'.join(
            f"{kp['name']}（错题 {kp['question_count']} 道，薄弱指数 {kp['weak_index']}）"
            for kp in recommendations[:3]
        )
        local_advice = f"今日建议优先复习：{names}。先从错题最多的知识点开始，做完后标记掌握情况。"
    else:
        local_advice = '当前学科还没有错题，先录入或导入错题，系统就会给出今日复习建议。'
    advice = local_advice
    if recommendations and db.get_setting('llm_api_key', ''):
        try:
            advice = call_llm(
                f'你是错题复习助手。根据这些薄弱知识点给出今天 3 条简短复习建议，每条不超过 30 字：\n{names}'
            ).strip() or local_advice
        except Exception:
            advice = local_advice
    return jsonify({
        'advice': advice,
        'recommendations': recommendations,
        'knowledge_points': all_kps
    })

@app.route('/api/review/plans', methods=['GET'])
def api_review_plans():
    return jsonify({'plans': db.list_review_plans(), 'counts': db.review_plan_counts()})

@app.route('/api/review/plans', methods=['POST'])
def api_add_review_plan():
    body = request.get_json(force=True, silent=True) or {}
    if not body.get('question_id') or not body.get('plan_date'):
        return jsonify({'error': '请选择错题和复习日期'}), 400
    return jsonify(db.add_review_plan(
        int(body['question_id']),
        body['plan_date'],
        body.get('note', '')
    ))

@app.route('/api/review/plans/generate', methods=['POST'])
def api_generate_review_plans():
    body = request.get_json(force=True, silent=True) or {}
    created = db.generate_today_plans(subject_id=body.get('subject_id'), limit=int(body.get('limit') or 10))
    return jsonify({'ok': True, 'created': len(created), 'plans': created})

@app.route('/api/review/plans/<int:pid>/complete', methods=['POST'])
def api_complete_review_plan(pid):
    body = request.get_json(force=True, silent=True) or {}
    ok = db.complete_review_plan(pid, int(body.get('result', 1)), body.get('note', ''))
    if not ok:
        return jsonify({'error': '计划不存在'}), 404
    return jsonify({'ok': True})

@app.route('/api/review/plans/<int:pid>', methods=['DELETE'])
def api_delete_review_plan(pid):
    db.delete_review_plan(pid)
    return jsonify({'ok': True})

@app.route('/api/kp_tree', methods=['GET'])
def api_kp_tree():
    subject_id = request.args.get('subject_id', type=int)
    if not subject_id:
        return jsonify({'error': 'subject_id required'}), 400
    return jsonify(db.kp_tree(subject_id))

# settings
@app.route('/api/settings', methods=['GET'])
def api_get_settings():
    return jsonify({
        'llm_api_key': db.get_setting('llm_api_key', ''),
        'llm_model': db.get_setting('llm_model', 'deepseek-chat'),
        'ocr_enabled': db.get_setting('ocr_enabled', 'true'),
        'ocr_lang': db.get_setting('ocr_lang', 'ch'),
        'ocr_use_orientation': db.get_setting('ocr_use_orientation', 'true'),
        'ocr_confidence_threshold': db.get_setting('ocr_confidence_threshold', '0.5'),
        'kp_auto_threshold': db.get_setting('kp_auto_threshold', '0.8'),
        'kp_pending_threshold': db.get_setting('kp_pending_threshold', '0.5')
    })

@app.route('/api/settings', methods=['POST'])
def api_save_settings():
    body = request.get_json(force=True, silent=True) or {}
    for k, v in body.items():
        db.set_setting(k, str(v))
    return jsonify({'ok': True})

# ---------- backup ----------
@app.route('/api/backup/export', methods=['GET'])
def api_export_package():
    uid = session.get('user_id')
    payload = json.dumps(db.export_user_data(uid), ensure_ascii=False, indent=2).encode('utf-8')
    sqlite_path = db.export_user_sqlite(uid)
    try:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
            zf.writestr('data.json', payload)
            zf.write(sqlite_path, 'data.db')
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/zip',
            as_attachment=True,
            download_name='wrongbook_backup.zip'
        )
    finally:
        try:
            os.remove(sqlite_path)
        except OSError:
            pass

@app.route('/api/backup/import', methods=['POST'])
def api_import_package():
    uid = session.get('user_id')
    if 'file' not in request.files:
        return jsonify({'error': '请选择备份数据包'}), 400
    raw = request.files['file'].read()
    try:
        if zipfile.is_zipfile(io.BytesIO(raw)):
            with zipfile.ZipFile(io.BytesIO(raw)) as zf:
                names = zf.namelist()
                if 'data.json' in names:
                    data = json.loads(zf.read('data.json').decode('utf-8'))
                elif 'data.db' in names:
                    temp_path = os.path.join(UPLOAD_DIR, f'import_{os.urandom(5).hex()}.db')
                    with open(temp_path, 'wb') as f:
                        f.write(zf.read('data.db'))
                    try:
                        data = db.read_sqlite_export(temp_path)
                    finally:
                        try:
                            os.remove(temp_path)
                        except OSError:
                            pass
                else:
                    return jsonify({'error': '数据包内容不正确'}), 400
        elif raw[:16].startswith(b'SQLite format 3'):
            temp_path = os.path.join(UPLOAD_DIR, f'import_{os.urandom(5).hex()}.db')
            with open(temp_path, 'wb') as f:
                f.write(raw)
            try:
                data = db.read_sqlite_export(temp_path)
            finally:
                try:
                    os.remove(temp_path)
                except OSError:
                    pass
        else:
            data = json.loads(raw.decode('utf-8'))
    except Exception as e:
        return jsonify({'error': f'备份解析失败：{e}'}), 400
    stats = db.import_user_data(uid, data)
    return jsonify({'ok': True, 'stats': stats})

@app.route('/api/backup/export/json', methods=['GET'])
def api_export_json():
    uid = session.get('user_id')
    data = json.dumps(db.export_user_data(uid), ensure_ascii=False, indent=2)
    return send_file(
        io.BytesIO(data.encode('utf-8')),
        mimetype='application/json',
        as_attachment=True,
        download_name='wrongbook_backup.json'
    )

@app.route('/api/backup/export/sqlite', methods=['GET'])
def api_export_sqlite():
    uid = session.get('user_id')
    path = db.export_user_sqlite(uid)
    return send_file(path, as_attachment=True, download_name='wrongbook_backup.db')

@app.route('/api/backup/import/json', methods=['POST'])
def api_import_json():
    uid = session.get('user_id')
    if 'file' not in request.files:
        return jsonify({'error': '请选择 JSON 备份文件'}), 400
    try:
        data = json.loads(request.files['file'].read().decode('utf-8'))
    except Exception as e:
        return jsonify({'error': f'文件解析失败：{e}'}), 400
    stats = db.import_user_data(uid, data)
    return jsonify({'ok': True, 'stats': stats})

@app.route('/api/backup/import/sqlite', methods=['POST'])
def api_import_sqlite():
    uid = session.get('user_id')
    if 'file' not in request.files:
        return jsonify({'error': '请选择 SQLite 备份文件'}), 400
    path = os.path.join(UPLOAD_DIR, f'import_{os.urandom(5).hex()}.db')
    request.files['file'].save(path)
    try:
        data = db.read_sqlite_export(path)
        stats = db.import_user_data(uid, data)
        return jsonify({'ok': True, 'stats': stats})
    finally:
        try:
            os.remove(path)
        except OSError:
            pass

# upload
@app.route('/api/upload_image', methods=['POST'])
def api_upload_image():
    if 'file' not in request.files:
        return jsonify({'error': 'no file'}), 400
    file = request.files['file']
    if not file or not file.filename:
        return jsonify({'error': 'no file'}), 400
    ext = os.path.splitext(file.filename)[1] or '.jpg'
    save_name = f"upload_{os.urandom(4).hex()}{ext}"
    save_path = os.path.join(UPLOAD_DIR, save_name)
    file.save(save_path)
    return jsonify({'url': f'/uploads/{save_name}', 'filename': save_name})

def ocr_image(image_path):
    ocr = get_ocr()
    # PaddleOCR 2.x 用 ocr()，3.x 用 predict()
    if hasattr(ocr, 'predict'):
        result = list(ocr.predict(image_path))
    else:
        result = ocr.ocr(image_path, cls=True)
    try:
        threshold = float(db.get_setting('ocr_confidence_threshold', '0.5') or 0.5)
    except (TypeError, ValueError):
        threshold = 0.5
    pairs = []

    def walk(node):
        if isinstance(node, (list, tuple)):
            if len(node) == 2 and isinstance(node[0], str) and isinstance(node[1], (int, float)):
                pairs.append((node[0], float(node[1])))
                return
            for child in node:
                walk(child)
        elif isinstance(node, dict):
            if 'rec_texts' in node and 'rec_scores' in node:
                for text, score in zip(node.get('rec_texts') or [], node.get('rec_scores') or []):
                    pairs.append((str(text), float(score)))
                return
            if 'text' in node:
                pairs.append((str(node['text']), float(node.get('score', 1))))
                return
            for value in node.values():
                walk(value)

    walk(result)
    # 去重并保留顺序
    seen = set()
    unique = []
    for text, score in pairs:
        if score < threshold or not text:
            continue
        if text not in seen:
            seen.add(text)
            unique.append(text)
    return '\n'.join(unique)

@app.route('/api/ocr', methods=['POST'])
def api_ocr():
    if 'file' not in request.files:
        return jsonify({'error': 'no file'}), 400
    file = request.files['file']
    if not file or not file.filename:
        return jsonify({'error': 'no file'}), 400
    ext = os.path.splitext(file.filename)[1] or '.jpg'
    save_name = f"ocr_{os.urandom(4).hex()}{ext}"
    save_path = os.path.join(UPLOAD_DIR, save_name)
    file.save(save_path)
    if not PADDLE_AVAILABLE:
        return jsonify({'error': 'PaddleOCR not installed', 'url': f'/uploads/{save_name}'}), 503
    try:
        text = ocr_image(save_path)
        return jsonify({
            'text': text,
            'url': f'/uploads/{save_name}'
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e), 'url': f'/uploads/{save_name}'}), 500

if __name__ == '__main__':
    db.init_db()
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
