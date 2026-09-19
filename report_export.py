#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""统计报表 Word / PDF 导出。"""
import io
import os
from datetime import datetime

import fitz
from docx import Document
from docx.shared import Pt

import db

FONT_PATH = 'C:/Windows/Fonts/simhei.ttf'


def collect_report(subject_id=None):
    questions = db.list_questions(subject_id=subject_id)
    status_counts = {0: 0, 1: 0, 2: 0, 3: 0}
    for q in questions:
        status_counts[q.get('status', 0)] = status_counts.get(q.get('status', 0), 0) + 1
    trend = db.trend_data(subject_id=subject_id, days=30)
    weak = [x for x in db.weak_analysis(subject_id=subject_id, top_n=10) if x.get('question_count')]
    plans = [p for p in db.list_review_plans() if p.get('status') == 'pending']
    subject_name = '全部学科'
    if subject_id:
        match = [s for s in db.list_subjects() if s['id'] == subject_id]
        if match:
            subject_name = match[0]['name']
    return {
        'subject_name': subject_name,
        'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M'),
        'totals': {
            'questions': len(questions),
            'knowledge_points': len(db.list_kps(subject_id)),
            'added_30d': sum(x['added'] for x in trend),
            'wrong_reviews_30d': sum(x['wrong_reviews'] for x in trend),
            'reviews': db.review_stats(subject_id),
            'plans': db.review_plan_counts(),
            'status_counts': status_counts
        },
        'weak': weak,
        'plans': plans,
        'trend': trend
    }


def build_word(data):
    doc = Document()
    doc.styles['Normal'].font.name = '微软雅黑'
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.add_heading('错题本知识溯源整理助手 · 统计报告', 0)
    doc.add_paragraph(f"学科：{data['subject_name']}")
    doc.add_paragraph(f"生成时间：{data['generated_at']}")

    doc.add_heading('一、总体统计', 1)
    t = data['totals']
    for text in [
        f"错题总数：{t['questions']}",
        f"知识点数量：{t['knowledge_points']}",
        f"近 30 天新增错题：{t['added_30d']}",
        f"近 30 天复习仍错：{t['wrong_reviews_30d']}",
        f"复习正确率：{t['reviews']['accuracy']}%",
        f"待复习计划：{t['plans']['total']}（今日 {t['plans']['today']}，逾期 {t['plans']['overdue']}）",
        "错题状态：" + "，".join(
            f"{name} {t['status_counts'].get(key, 0)}"
            for key, name in [(0, '未掌握'), (1, '已掌握'), (2, '已复习'), (3, '已归档')]
        )
    ]:
        doc.add_paragraph(text)

    doc.add_heading('二、薄弱知识点 TOP-N', 1)
    if data['weak']:
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        headers = ['知识点', '错题数', '累计错误', '近7天', '薄弱指数']
        for i, header in enumerate(headers):
            table.rows[0].cells[i].text = header
        for item in data['weak']:
            cells = table.add_row().cells
            cells[0].text = item['name']
            cells[1].text = str(item['question_count'])
            cells[2].text = str(item['error_count'])
            cells[3].text = str(item['recent_errors'])
            cells[4].text = str(item['weak_index'])
    else:
        doc.add_paragraph('暂无可分析的薄弱知识点。')

    doc.add_heading('三、复习计划', 1)
    if data['plans']:
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        for i, header in enumerate(['日期', '错题', '备注']):
            table.rows[0].cells[i].text = header
        for plan in data['plans']:
            cells = table.add_row().cells
            cells[0].text = plan.get('plan_date', '')
            cells[1].text = plan.get('question_text', '')
            cells[2].text = plan.get('note', '')
    else:
        doc.add_paragraph('暂无待复习计划。')

    doc.add_heading('四、近 30 天趋势', 1)
    trend_table = doc.add_table(rows=1, cols=3)
    trend_table.style = 'Table Grid'
    for i, header in enumerate(['日期', '新增错题', '复习仍错']):
        trend_table.rows[0].cells[i].text = header
    for item in data['trend']:
        cells = trend_table.add_row().cells
        cells[0].text = item['date']
        cells[1].text = str(item['added'])
        cells[2].text = str(item['wrong_reviews'])

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def build_pdf(data):
    doc = fitz.open()
    page = doc.new_page()
    fontname = 'china'
    fontfile = FONT_PATH if os.path.exists(FONT_PATH) else None
    if fontfile:
        page.insert_font(fontname=fontname, fontfile=fontfile)
    y = 55

    def ensure_space(height=20):
        nonlocal page, y
        if y + height > 800:
            page = doc.new_page()
            if fontfile:
                page.insert_font(fontname=fontname, fontfile=fontfile)
            y = 55

    def write(text, size=11, gap=18, color=(0.1, 0.1, 0.1)):
        nonlocal y
        ensure_space(gap)
        page.insert_text((48, y), str(text), fontsize=size, fontname=fontname, color=color)
        y += gap

    write('错题本知识溯源整理助手 · 统计报告', size=18, gap=30, color=(0.12, 0.35, 0.75))
    write(f"学科：{data['subject_name']}")
    write(f"生成时间：{data['generated_at']}", gap=24)
    t = data['totals']
    write('一、总体统计', size=14, gap=22, color=(0.05, 0.45, 0.42))
    for text in [
        f"错题总数：{t['questions']}",
        f"知识点数量：{t['knowledge_points']}",
        f"近30天新增错题：{t['added_30d']}",
        f"近30天复习仍错：{t['wrong_reviews_30d']}",
        f"复习正确率：{t['reviews']['accuracy']}%",
        f"待复习计划：{t['plans']['total']}（今日 {t['plans']['today']}，逾期 {t['plans']['overdue']}）"
    ]:
        write(text)

    y += 8
    write('二、薄弱知识点 TOP-N', size=14, gap=22, color=(0.05, 0.45, 0.42))
    write('知识点                 错题数   累计错误   近7天   薄弱指数', size=10, gap=18)
    for item in data['weak']:
        write(f"{item['name'][:14]:<16} {item['question_count']:>6} {item['error_count']:>9} {item['recent_errors']:>7} {item['weak_index']:>9}", size=10)
    if not data['weak']:
        write('暂无可分析的薄弱知识点。')

    y += 8
    write('三、待复习计划', size=14, gap=22, color=(0.05, 0.45, 0.42))
    for plan in data['plans'][:30]:
        write(f"{plan.get('plan_date','')}  {plan.get('question_text','')[:36]}", size=10)
    if not data['plans']:
        write('暂无待复习计划。')

    if fontfile and hasattr(doc, 'subset_fonts'):
        doc.subset_fonts()
    out = io.BytesIO()
    doc.save(out)
    doc.close()
    out.seek(0)
    return out
